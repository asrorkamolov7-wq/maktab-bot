"""Excel hisobot (openpyxl). Sinxron funksiya: asyncio.to_thread orqali chaqiriladi."""
from __future__ import annotations

import re
from datetime import datetime, timedelta, timezone
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from models import Survey
from utils import grade_key

UZ_TZ = timezone(timedelta(hours=5))  # O'zbekistonda yoz/qish vaqti yo'q

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
ZEBRA_FILL = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
_SIDE = Side(style="thin", color="D9D9D9")
BORDER = Border(left=_SIDE, right=_SIDE, top=_SIDE, bottom=_SIDE)

_BAD_SHEET_CHARS = re.compile(r"[\\/*?:\[\]]")


def _fmt_dt(value: str) -> str:
    try:
        dt = datetime.strptime(value, "%Y-%m-%d %H:%M:%S").replace(tzinfo=timezone.utc)
        return dt.astimezone(UZ_TZ).strftime("%d.%m.%Y %H:%M")
    except (ValueError, TypeError):
        return value or ""


def _set(ws, row: int, col: int, value: Any):
    """Xavfsiz yozish: '=...' formula sifatida talqin qilinmasin, noqonuniy belgilar olib tashlansin."""
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)
    cell = ws.cell(row=row, column=col, value=value)
    if isinstance(value, str):
        cell.data_type = "s"
    return cell


def _sheet_name(survey: Survey, used: set[str]) -> str:
    g = survey.grades
    base = f"{g[0]}-{g[-1]}-sinf" if len(g) > 1 else f"{g[0]}-sinf"
    base = _BAD_SHEET_CHARS.sub("", base)[:28]
    name, n = base, 2
    while name in used:
        name = f"{base} ({n})"
        n += 1
    used.add(name)
    return name


def _style_header(ws, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 45


def build_report(data: list[dict[str, Any]]) -> bytes:
    wb = Workbook()
    summary = wb.active
    summary.title = "Umumiy"
    _set(summary, 1, 1, "So'rovnoma")
    _set(summary, 1, 2, "Sinf")
    _set(summary, 1, 3, "Javoblar soni")
    _style_header(summary, 3)
    summary.row_dimensions[1].height = 25

    used_names = {"Umumiy"}
    srow = 2

    for item in data:
        survey: Survey = item["survey"]
        subs = sorted(item["submissions"], key=lambda s: (grade_key(s["grade"]), s["full_name"].lower()))
        ws = wb.create_sheet(_sheet_name(survey, used_names))

        base_headers = ["T/r", "O'quvchining F.I.SH", "Sinf", "Telegram ID", "Username", "Sana"]
        headers = base_headers + [f"{i}. {q.text}" for i, q in enumerate(survey.questions, start=1)]
        for c, h in enumerate(headers, start=1):
            _set(ws, 1, c, h)
        _style_header(ws, len(headers))

        for r, s in enumerate(subs, start=2):
            username = f"@{s['username']}" if s["username"] else "-"
            values = [r - 1, s["full_name"], s["grade"], s["user_id"], username, _fmt_dt(s["created_at"])]
            values += [a if a else "-" for a in s["answers"]]
            for c, v in enumerate(values, start=1):
                cell = _set(ws, r, c, v)
                cell.border = BORDER
                if r % 2 == 1:
                    cell.fill = ZEBRA_FILL
                center = c in (1, 3, 4, 6)
                cell.alignment = Alignment(
                    horizontal="center" if center else "left", vertical="center", wrap_text=not center
                )

        widths = {1: 6, 2: 30, 3: 8, 4: 14, 5: 18, 6: 17}
        for c in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(c)].width = widths.get(c, 32)
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(subs) + 1}"

        # Umumiy sahifa: sinflar bo'yicha sonlar
        counts: dict[str, int] = {}
        for s in subs:
            counts[s["grade"]] = counts.get(s["grade"], 0) + 1
        first = True
        for grade in sorted(counts, key=grade_key):
            _set(summary, srow, 1, survey.title if first else "")
            _set(summary, srow, 2, grade)
            _set(summary, srow, 3, counts[grade])
            for c in (1, 2, 3):
                summary.cell(row=srow, column=c).border = BORDER
                summary.cell(row=srow, column=c).alignment = Alignment(
                    horizontal="left" if c == 1 else "center", vertical="center", wrap_text=True
                )
            first = False
            srow += 1
        _set(summary, srow, 1, "")
        _set(summary, srow, 2, "Jami")
        _set(summary, srow, 3, len(subs))
        for c in (1, 2, 3):
            cell = summary.cell(row=srow, column=c)
            cell.font = Font(bold=True)
            cell.border = BORDER
            cell.fill = ZEBRA_FILL
            cell.alignment = Alignment(horizontal="center")
        srow += 2

    summary.column_dimensions["A"].width = 60
    summary.column_dimensions["B"].width = 12
    summary.column_dimensions["C"].width = 16

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
