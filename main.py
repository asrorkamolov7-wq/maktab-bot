"""Maktab maslahatchisi boti — bitta faylli versiya (aiogram 3).

Muhit o'zgaruvchilari: BOT_TOKEN, ADMIN_IDS (majburiy); DB_PATH, PORT, GRADE_LETTERS (ixtiyoriy).
Ishga tushirish: python main.py
"""
from __future__ import annotations

import aiosqlite
import asyncio
import json
import logging
import os
import re
import time
import uuid

from dataclasses import (
    dataclass,
    field,
)
from typing import (
    Any,
    Awaitable,
    Callable,
    TYPE_CHECKING,
)
from datetime import (
    datetime,
    timedelta,
    timezone,
)
from io import BytesIO
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter
from html import escape
from aiogram.types import (
    BotCommand,
    BotCommandScopeChat,
    BufferedInputFile,
    CallbackQuery,
    ErrorEvent,
    InlineKeyboardButton,
    InlineKeyboardMarkup,
    KeyboardButton,
    Message,
    ReplyKeyboardMarkup,
    TelegramObject,
    User,
)
from aiogram.fsm.state import (
    State,
    StatesGroup,
)
from aiogram.filters import (
    BaseFilter,
    Command,
    CommandStart,
    StateFilter,
)
from aiogram import (
    BaseMiddleware,
    Bot,
    Dispatcher,
    F,
    Router,
)
from aiogram.exceptions import (
    SkipHandler,
    TelegramAPIError,
    TelegramBadRequest,
    TelegramForbiddenError,
    TelegramRetryAfter,
)
from aiogram.fsm.context import FSMContext
from aiohttp import web
from aiogram.client.default import DefaultBotProperties
from aiogram.enums import ParseMode
from aiogram.fsm.storage.memory import (
    MemoryStorage,
    SimpleEventIsolation,
)


# ====================================================================
# CONFIG
# ====================================================================
@dataclass(frozen=True)
class Settings:
    bot_token: str
    admin_ids: frozenset[int]
    db_path: str
    port: int
    grade_letters: tuple[str, ...]
    broadcast_delay: float = 0.05  # Telegram limiti: ~30 xabar/soniya


def _parse_ids(raw: str) -> frozenset[int]:
    ids: set[int] = set()
    for part in raw.replace(";", ",").split(","):
        part = part.strip()
        if not part:
            continue
        try:
            ids.add(int(part))
        except ValueError:
            raise RuntimeError(f"ADMIN_IDS ichida noto'g'ri qiymat: {part!r}") from None
    return frozenset(ids)


def load_settings() -> Settings:
    token = os.getenv("BOT_TOKEN", "").strip()
    if not token:
        raise RuntimeError("BOT_TOKEN muhit o'zgaruvchisi o'rnatilmagan.")

    # ADMIN_ID (eski nom) ham qabul qilinadi
    admin_ids = _parse_ids(os.getenv("ADMIN_IDS") or os.getenv("ADMIN_ID", ""))
    if not admin_ids:
        raise RuntimeError("ADMIN_IDS muhit o'zgaruvchisi o'rnatilmagan.")

    letters = tuple(
        x.strip().upper()
        for x in os.getenv("GRADE_LETTERS", "A,B,V,G,D").split(",")
        if x.strip()
    )
    return Settings(
        bot_token=token,
        admin_ids=admin_ids,
        db_path=os.getenv("DB_PATH", "data/school_bot.db"),
        port=int(os.getenv("PORT", "10000")),
        grade_letters=letters or ("A", "B", "V", "G", "D"),
    )


# ====================================================================
# MODELS
# ====================================================================
_NUMBERING_RE = re.compile(r"^\s*\d+\s*[.)]\s*")


def clean_question_text(text: str) -> str:
    """'1. Savol' -> 'Savol' (raqamlash botning o'zida qo'yiladi)."""
    return _NUMBERING_RE.sub("", text).strip()


@dataclass
class Question:
    text: str
    type: str = "text"  # "text" yoki "choice"
    options: list[str] = field(default_factory=list)

    @classmethod
    def from_dict(cls, d: dict) -> "Question":
        return cls(
            text=clean_question_text(d.get("question") or d.get("text") or ""),
            type=d.get("type", "text"),
            options=list(d.get("options") or []),
        )

    def to_dict(self) -> dict:
        return {"type": self.type, "question": self.text, "options": self.options}


@dataclass
class Survey:
    id: str
    title: str
    grades: list[int]
    questions: list[Question]
    is_active: bool = True
    created_at: str = ""


# ====================================================================
# TEXTS
# ====================================================================
ABOUT_COUNSELOR = """
👨‍💼 <b>Maktab maslahatchisi kabineti</b>
━━━━━━━━━━━━━━━━━━━━
• <b>F.I.SH:</b> Kamolov Asror
• <b>Lavozimi:</b> Maktab maslahatchisi
• <b>Xona:</b> 2-qavat, Maslahatchi xonasi
• <b>Qabul vaqti:</b> 09:00 dan 16:00 gacha
• <b>Bog'lanish uchun telefon:</b> +998 33 443 55 77

<i>Har qanday shaxsiy masala, kasb tanlash yoki qiyinchiliklar bo'yicha bemalol murojaat qilishingiz mumkin. Barcha suhbatlar to'liq sir saqlanadi.</i>
"""

CLUBS_INFO = """
🎭 <b>Maktabda faoliyat yuritayotgan 7 ta rivojlanish klubi:</b>
━━━━━━━━━━━━━━━━━━━━
1. 🎭 <b>Turon teatr</b> — Sahna madaniyati, aktyorlik mahorati va ijodiy qobiliyatlarni rivojlantirish.
2. 🎵 <b>Iqtidor ansambli</b> — Musiqa, milliy va zamonaviy ijrochilik sanʼati.
3. 📜 <b>Jadidlar izidan</b> — Tarixiy meros, kitobxonlik va maʼrifatparvarlik tadbirlari.
4. 🌿 <b>Eco-schools Uzbekistan</b> — Ekologik savodxonlik, tabiatni asrash va atrof-muhit loyihalari.
5. 🌍 <b>Xorijiy tillar</b> — Chet tillarida erkin muloqot va soʻzlashuv klubi.
6. 🎤 <b>Debat</b> — Notiqlik, tanqidiy fikrlash va munozara madaniyati.
7. 💻 <b>Raqamli avlod qizlari</b> — Qizlar uchun IT, zamonaviy kasblar va media savodxonlik.

<i>Klublarga a'zo bo'lish uchun maktab maslahatchisi xonasiga (2-qavat) murojaat qiling.</i>
"""

COUNCIL_INFO = """
👥 <b>O'quvchilar kengashi tashkiliy tuzilmasi</b>
━━━━━━━━━━━━━━━━━━━━
👑 <b>Kelajak lideri</b>
O'quvchilar kengashi faoliyatiga umumiy rahbarlik qiladi.

📌 <b>Asosiy yo'nalish koordinatorlari va vazifalari:</b>

🎓 <b>Ta'lim koordinatori</b>
Ta'lim sifatini oshirish, fan va til to'garaklarini tashkil etish, o'quvchilarning o'zlashtirishiga yordam berish.

🎭 <b>Madaniyat koordinatori</b>
Maktabning madaniy tadbirlarini tashkil etish, milliy qadriyatlarni targ'ib qilish, "Turon maktab teatri" va "Iqtidor ansambli" faoliyatini qo'llab-quvvatlash.

💻 <b>IT va sun'iy intellekt koordinatori</b>
Raqamli texnologiyalarni targ'ib qilish, IT ko'nikmalarini rivojlantirish, innovatsion loyihalar va tanlovlarni tashkil etish.

🏃‍♂️ <b>Sport koordinatori</b>
Sport tadbirlarini tashkil etish, sog'lom turmush tarzini targ'ib qilish, sport klublari va tanlovlarni rivojlantirish.

🌿 <b>Ekologiya koordinatori</b>
Ekologik loyihalar va aksiyalarni tashkil etish, tabiatni muhofaza qilish, ekologik madaniyatni targ'ib qilish.

📢 <b>Media koordinatori</b>
Maktab hayotini ommaviy axborotda yoritish, ijtimoiy tarmoqlarni yuritish, media loyihalarini tashkil etish.

👥 <b>Ijtimoiy tashabbuslar koordinatori</b>
O'quvchilar tashabbuslarini qo'llab-quvvatlash, ijtimoiy loyihalar va volontyorlik faoliyatini tashkil etish, inklyuziv muhitni rivojlantirish.

━━━━━━━━━━━━━━━━━━━━
👥 <b>5–11-sinf vakillari:</b> Sinf o'quvchilari tomonidan saylanadi.
☘️ <b>Klublar:</b> Maktabda faoliyat yurituvchi tarmoq to'garaklari.
"""

WELCOME = (
    "Assalomu alaykum, <b>{name}</b>!\n\n"
    "🏫 <b>Maktab maslahatchisining rasmiy portaliga xush kelibsiz.</b>\n"
    "Ushbu bot orqali siz o'z qiziqishlaringiz bo'yicha so'rovnomalarda qatnashishingiz, "
    "maktab klublari va kengash faoliyati bilan tanishishingiz hamda "
    "maslahatchiga to'g'ridan-to'g'ri murojaat yo'llashingiz mumkin.\n\n"
    "Kerakli bo'limni tanlang:"
)

SURVEY_TEMPLATE = (
    "Sarlavha: Bitiruvchilar rejalari\n"
    "Sinflar: 9-11\n\n"
    "1. Qaysi sohaga qiziqasiz?\n"
    "- Axborot texnologiyalari (IT)\n"
    "- Tibbiyot\n"
    "- Iqtisodiyot va biznes\n\n"
    "2. Qaysi oliygohga topshirmoqchisiz?\n"
    "[yozma]"
)


def _c(question: str, options: list[str]) -> dict:
    return {"type": "choice", "question": question, "options": options}


def _t(question: str) -> dict:
    return {"type": "text", "question": question, "options": []}


DEFAULT_SURVEYS = [
    {
        "id": "survey_5_8",
        "title": "«Mening qiziqishlarim va kelajakdagi kasbim» (5–8-sinf)",
        "grades": [5, 6, 7, 8],
        "questions": [
            _c("Maktabda qaysi fanlarni o'rganish sizga ko'proq yoqadi?", ["Matematika", "Ona tili va adabiyot", "Tarix", "Chet tili", "Informatika", "Biologiya", "Fizika", "Jismoniy tarbiya", "Musiqa / Tasviriy san'at"]),
            _c("Bo'sh vaqtingizda nimalar bilan shug'ullanishni yoqtirasiz?", ["Kitob o'qish", "Sport", "Rasm chizish", "Kompyuter va IT", "Video / montaj / dizayn", "Qo'shiq / cholg'u", "Pazandachilik / hunar"]),
            _c("O'zingizda qaysi qobiliyat kuchliroq deb o'ylaysiz?", ["Hisoblash va mantiq", "Chiroyli gapirish", "Texnika va kompyuter", "Musiqa va ijod", "Sport", "Tashkilotchilik", "Tinglash va yordam berish", "Hali bilmayman"]),
            _c("Yangi narsani o'rganishda qaysi usul sizga ko'proq yoqadi?", ["O'qib o'rganish", "Amalda bajarib ko'rish", "Video orqali ko'rish", "Ustozdan so'rab", "Do'stlar bilan birgalikda", "Tajriba qilib ko'rish"]),
            _c("Quyidagi faoliyatlardan qaysi biri sizga eng qiziq?", ["Kompyuter / IT", "Tibbiyot", "Texnika va muhandislik", "San'at va ijod", "Sport", "Ta'lim va pedagogika", "Biznes va tadbirkorlik", "Huquq / Harbiy"]),
            _t("Kelajakda qaysi kasb yoki kasblarga qiziqasiz? Nomlarini yozing:"),
            _c("Bu kasbga qiziqishingizga nima sabab bo'lgan?", ["Bu kasb menga yoqadi", "Tanishim shu sohada ishlaydi", "Bolalikdagi orzuim", "Daromadi yaxshi", "Odamlarga yordam berish", "Hali aniq bilmayman"]),
            _c("Siz ko'proq qanday ishni bajarishni xohlaysiz?", ["Odamlar bilan ishlash", "Kompyuter bilan ishlash", "Texnika va qurilmalar", "Hayvonlar va tabiat", "Ijodiy ishlar", "Jamoa bilan ishlash", "Mustaqil ishlash"]),
            _c("Qaysi to'garak yoki klub sizni ko'proq qiziqtiradi?", ["Turon teatr", "Iqtidor ansambli", "Jadidlar izidan", "Eco-schools Uzbekistan", "Xorijiy tillar", "Debat klubi", "Raqamli avlod qizlari"]),
            _t("Maktabda yangi to'garak ochilsa, qaysi yo'nalishda bo'lishini xohlardingiz?"),
            _c("O'zingizni kelajakda qanday inson sifatida tasavvur qilasiz?", ["Yaxshi mutaxassis", "Tadbirkor", "Olim / Tadqiqotchi", "Ijodkor", "Sportchi", "Rahbar", "O'qituvchi"]),
            _c("Sizga kasb tanlashda kimning maslahati ko'proq ta'sir qiladi?", ["Ota-onam", "O'qituvchim", "Sinfdoshlarim / do'stlarim", "Maktab maslahatchisi", "O'zim mustaqil tanlayman"]),
            _t("Kelajakdagi qaysi kasb haqida ko'proq ma'lumot olishni xohlaysiz?"),
            _t("Siz maktab hayotini yaxshilash uchun qanday yangi taklif bera olasiz?"),
            _t("O'zingiz haqingizda biz bilishimiz kerak bo'lgan boshqa qiziqishingiz bormi?"),
        ],
    },
    {
        "id": "survey_9_11",
        "title": "«Kasb tanlash va oliygohga tayyorgarlik» (9–11-sinf)",
        "grades": [9, 10, 11],
        "questions": [
            _c("Kelajakda qaysi kasb egasi bo'lishni xohlaysiz?", ["Tibbiyot", "Muhandislik / Arxitektura", "IT sohasi", "Iqtisodiyot / Biznes", "Ta'lim / Pedagogika", "Huquqshunoslik", "Harbiy soha", "Boshqa"]),
            _c("Sizni qaysi fanlar ko'proq qiziqtiradi?", ["Matematika", "Fizika", "Kimyo / Biologiya", "Ingliz tili / Chet tillari", "Tarix va huquq", "Ona tili va adabiyot"]),
            _c("Maktabni tugatgach o'qishni davom ettirishni rejalashtiryapsizmi?", ["Ha, albatta", "Yo'q", "Hali aniq emas"]),
            _c("Qayerda o'qishni xohlaysiz?", ["O'zbekiston oliy ta'lim muassasalarida", "Xorijiy davlat universitetlarida", "Ikkalasini ham ko'rib chiqyapman"]),
            _c("Qaysi davlat universitetlarida o'qish sizni qiziqtiradi?", ["O'zbekiston", "AQSh", "Buyuk Britaniya", "Koreya", "Turkiya", "Yevropa davlatlari", "Rossiya"]),
            _c("Xorijiy universitetlarga hujjat topshirishni rejalashtiryapsizmi?", ["Ha", "Yo'q", "Hali o'ylab ko'ryapman"]),
            _c("Qaysi xalqaro test yoki sertifikatlarga tayyorlanyapsiz?", ["IELTS", "SAT", "CEFR", "TOEFL", "Hali boshlamaganman"]),
            _c("Ingliz tilini bilish darajangiz qanday?", ["Boshlang'ich", "O'rta", "Yaxshi (B2 daraja)", "Juda yaxshi (C1 daraja)"]),
            _t("Hozirda qaysi to'garak yoki o'quv markazlariga qatnashyapsiz?"),
            _c("Siz qaysi sohada o'zingizni rivojlantirishni xohlaysiz?", ["Akademik bilimlar", "Liderlik ko'nikmalari", "Ilmiy izlanish", "Sport", "Ijodiy yo'nalishlar"]),
            _c("Universitetga kirish uchun sizga birinchi navbatda qanday yordam kerak?", ["Testlarga tayyorgarlik", "Universitet tanlashda maslahat", "Hujjat topshirish tartibi", "Grant imkoniyatlari"]),
            _t("Kelajakdagi asosiy maqsadingiz nima? Qisqacha yozing:"),
            _c("Agar chet elda o'qish imkoniyati yaratilsa, undan foydalanishga tayyormisiz?", ["Ha, albatta", "Ehtimol", "Yo'q"]),
            _c("Maktab maslahatchisi bilan suhbatlashishni xohlaysizmi?", ["Ha, suhbatlashmoqchiman", "Hozircha yo'q"]),
            _t("Maktab hayotini yaxshilash uchun qanday taklifingiz bor?"),
        ],
    },
]


# ====================================================================
# SURVEY_PARSER
# ====================================================================
MAX_QUESTIONS = 50
MAX_OPTIONS = 12
MAX_OPTION_LEN = 60
MIN_GRADE, MAX_GRADE = 5, 11

_QUESTION_RE = re.compile(r"^\d+\s*[.)]\s*(.+)$")
_GRADES_RE = re.compile(r"(\d+)(?:\s*[-–—]\s*(\d+))?")


def parse_grades(text: str) -> list[int]:
    """'9, 10, 11' yoki '5-8' yoki '5-6, 9' -> [..]"""
    grades: set[int] = set()
    for start, end in _GRADES_RE.findall(text):
        a = int(start)
        b = int(end) if end else a
        if a > b:
            a, b = b, a
        grades.update(range(a, b + 1))
    return sorted(grades)


def parse_survey_text(raw: str) -> tuple[str, list[int], list[Question], list[str]]:
    title = ""
    grades: list[int] = []
    questions: list[Question] = []
    errors: list[str] = []
    current: Question | None = None
    explicit_text = False  # joriy savolda [yozma] belgisi bormi

    def close_current() -> None:
        nonlocal current, explicit_text
        if current is None:
            return
        if explicit_text and current.options:
            errors.append(f"«{current.text[:40]}»: variantlar ham, [yozma] ham berilgan.")
        elif current.options:
            if len(current.options) < 2:
                errors.append(f"«{current.text[:40]}»: kamida 2 ta variant kerak.")
            if len(current.options) > MAX_OPTIONS:
                errors.append(f"«{current.text[:40]}»: variantlar {MAX_OPTIONS} tadan oshmasin.")
            current.type = "choice"
        else:
            current.type = "text"
        questions.append(current)
        current = None
        explicit_text = False

    for line in (ln.strip() for ln in raw.splitlines()):
        if not line:
            continue
        low = line.lower()
        if low.startswith(("sarlavha:", "title:")):
            title = line.split(":", 1)[1].strip()
        elif low.startswith(("sinflar:", "sinf:", "grades:")):
            grades = parse_grades(line.split(":", 1)[1])
        elif (m := _QUESTION_RE.match(line)):
            close_current()
            current = Question(text=m.group(1).strip())
        elif line[0] in "-•–":
            if current is None:
                errors.append(f"Variant savoldan oldin yozilgan: «{line[:40]}»")
                continue
            opt = line.lstrip("-•– ").strip()
            if not opt:
                continue
            if len(opt) > MAX_OPTION_LEN:
                errors.append(f"Variant juda uzun (>{MAX_OPTION_LEN}): «{opt[:30]}…»")
                continue
            current.options.append(opt)
        elif low == "[yozma]":
            if current is None:
                errors.append("[yozma] savoldan oldin yozilgan.")
            else:
                explicit_text = True
        else:
            errors.append(f"Tushunarsiz qator: «{line[:50]}»")
    close_current()

    if not title:
        errors.append("«Sarlavha:» qatori topilmadi.")
    elif len(title) > 200:
        errors.append("Sarlavha 200 belgidan oshmasin.")
    if not grades:
        errors.append("«Sinflar:» qatori topilmadi yoki bo'sh (masalan: Sinflar: 9-11).")
    elif any(g < MIN_GRADE or g > MAX_GRADE for g in grades):
        errors.append(f"Sinflar faqat {MIN_GRADE} dan {MAX_GRADE} gacha bo'lishi mumkin.")
    if not questions:
        errors.append("Kamida 1 ta savol kerak (masalan: «1. Savol matni»).")
    elif len(questions) > MAX_QUESTIONS:
        errors.append(f"Savollar {MAX_QUESTIONS} tadan oshmasin.")

    return title, grades, questions, errors


# ====================================================================
# DATABASE
# ====================================================================
logger = logging.getLogger(__name__)

SCHEMA = """
CREATE TABLE IF NOT EXISTS meta (
    key   TEXT PRIMARY KEY,
    value TEXT
);

CREATE TABLE IF NOT EXISTS users (
    user_id    INTEGER PRIMARY KEY,
    full_name  TEXT,
    username   TEXT,
    is_blocked INTEGER NOT NULL DEFAULT 0,
    joined_at  TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
    last_seen  TIMESTAMP DEFAULT CURRENT_TIMESTAMP
);

CREATE TABLE IF NOT EXISTS surveys (
    id             TEXT PRIMARY KEY,
    title          TEXT NOT NULL,
    grades_json    TEXT NOT NULL,
    questions_json TEXT NOT NULL,
    is_active      INTEGER NOT NULL DEFAULT 1,
    created_at     TIMESTAMP DEFAULT CURRENT_TIMESTAMP
);

CREATE TABLE IF NOT EXISTS submissions (
    id         INTEGER PRIMARY KEY AUTOINCREMENT,
    survey_id  TEXT NOT NULL REFERENCES surveys(id),
    user_id    INTEGER NOT NULL,
    full_name  TEXT NOT NULL,
    grade      TEXT NOT NULL,
    created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
    UNIQUE (survey_id, user_id)
);

CREATE TABLE IF NOT EXISTS submission_answers (
    submission_id INTEGER NOT NULL REFERENCES submissions(id) ON DELETE CASCADE,
    q_idx         INTEGER NOT NULL,
    question      TEXT NOT NULL,
    answer        TEXT NOT NULL,
    PRIMARY KEY (submission_id, q_idx)
);

CREATE TABLE IF NOT EXISTS feedback_map (
    admin_id     INTEGER NOT NULL,
    admin_msg_id INTEGER NOT NULL,
    user_id      INTEGER NOT NULL,
    user_name    TEXT,
    created_at   TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
    PRIMARY KEY (admin_id, admin_msg_id)
);

CREATE INDEX IF NOT EXISTS idx_submissions_user ON submissions(user_id);
"""


def _row_to_survey(row: aiosqlite.Row) -> Survey:
    return Survey(
        id=row["id"],
        title=row["title"],
        grades=json.loads(row["grades_json"]),
        questions=[Question.from_dict(q) for q in json.loads(row["questions_json"])],
        is_active=bool(row["is_active"]),
        created_at=row["created_at"] or "",
    )


class Database:
    def __init__(self, conn: aiosqlite.Connection) -> None:
        self._conn = conn
        self._write_lock = asyncio.Lock()

    # ---------- hayot sikli ----------
    @classmethod
    async def connect(cls, path: str) -> "Database":
        folder = os.path.dirname(os.path.abspath(path))
        os.makedirs(folder, exist_ok=True)
        conn = await aiosqlite.connect(path)
        conn.row_factory = aiosqlite.Row
        await conn.execute("PRAGMA journal_mode=WAL")
        await conn.execute("PRAGMA synchronous=NORMAL")
        await conn.execute("PRAGMA foreign_keys=ON")
        db = cls(conn)
        await db._init_schema()
        return db

    async def close(self) -> None:
        await self._conn.close()

    async def _init_schema(self) -> None:
        async with self._write_lock:
            await self._conn.executescript(SCHEMA)
            await self._migrate_legacy()
            await self._seed_defaults()
            await self._conn.commit()

    async def _migrate_legacy(self) -> None:
        """Eski versiyadagi `custom_surveys` jadvalidagi so'rovnomalarni ko'chirish."""
        async with self._conn.execute(
            "SELECT name FROM sqlite_master WHERE type='table' AND name='custom_surveys'"
        ) as cur:
            if await cur.fetchone() is None:
                return
        await self._conn.execute(
            """INSERT OR IGNORE INTO surveys (id, title, grades_json, questions_json, created_at)
               SELECT id, title, grades_json, questions_json, created_at FROM custom_surveys"""
        )
        await self._conn.execute(
            "INSERT OR IGNORE INTO meta (key, value) VALUES ('defaults_seeded', '1')"
        )
        logger.info("Eski so'rovnomalar yangi jadvalga ko'chirildi.")

    async def _seed_defaults(self) -> None:
        """Boshlang'ich so'rovnomalar faqat BIR marta qo'shiladi (o'chirilsa qaytmaydi)."""
        async with self._conn.execute(
            "SELECT value FROM meta WHERE key='defaults_seeded'"
        ) as cur:
            if await cur.fetchone():
                return
        for s in DEFAULT_SURVEYS:
            await self._conn.execute(
                "INSERT OR IGNORE INTO surveys (id, title, grades_json, questions_json) VALUES (?,?,?,?)",
                (s["id"], s["title"], json.dumps(s["grades"]), json.dumps(s["questions"], ensure_ascii=False)),
            )
        await self._conn.execute("INSERT INTO meta (key, value) VALUES ('defaults_seeded', '1')")

    async def _execute(self, sql: str, params: tuple | list = ()) -> int:
        async with self._write_lock:
            cur = await self._conn.execute(sql, params)
            await self._conn.commit()
            return cur.rowcount

    async def _fetchall(self, sql: str, params: tuple | list = ()) -> list[aiosqlite.Row]:
        async with self._conn.execute(sql, params) as cur:
            return list(await cur.fetchall())

    async def _fetchone(self, sql: str, params: tuple | list = ()) -> aiosqlite.Row | None:
        async with self._conn.execute(sql, params) as cur:
            return await cur.fetchone()

    # ---------- foydalanuvchilar ----------
    async def upsert_user(self, user_id: int, full_name: str, username: str | None) -> None:
        await self._execute(
            """INSERT INTO users (user_id, full_name, username) VALUES (?,?,?)
               ON CONFLICT(user_id) DO UPDATE SET
                   full_name = excluded.full_name,
                   username  = excluded.username,
                   is_blocked = 0,
                   last_seen = CURRENT_TIMESTAMP""",
            (user_id, full_name, username or ""),
        )

    async def get_broadcast_ids(self) -> list[int]:
        rows = await self._fetchall("SELECT user_id FROM users WHERE is_blocked = 0")
        return [r["user_id"] for r in rows]

    async def mark_blocked(self, user_id: int) -> None:
        await self._execute("UPDATE users SET is_blocked = 1 WHERE user_id = ?", (user_id,))

    async def user_exists(self, user_id: int) -> bool:
        return await self._fetchone("SELECT 1 FROM users WHERE user_id = ?", (user_id,)) is not None

    # ---------- so'rovnomalar ----------
    async def list_surveys(self, only_active: bool = False) -> list[Survey]:
        sql = "SELECT * FROM surveys"
        if only_active:
            sql += " WHERE is_active = 1"
        sql += " ORDER BY created_at, rowid"
        return [_row_to_survey(r) for r in await self._fetchall(sql)]

    async def get_survey(self, survey_id: str) -> Survey | None:
        row = await self._fetchone("SELECT * FROM surveys WHERE id = ?", (survey_id,))
        return _row_to_survey(row) if row else None

    async def add_survey(self, survey_id: str, title: str, grades: list[int], questions: list[Question]) -> None:
        await self._execute(
            "INSERT INTO surveys (id, title, grades_json, questions_json) VALUES (?,?,?,?)",
            (
                survey_id,
                title,
                json.dumps(grades),
                json.dumps([q.to_dict() for q in questions], ensure_ascii=False),
            ),
        )

    async def set_survey_active(self, survey_id: str, active: bool) -> None:
        await self._execute("UPDATE surveys SET is_active = ? WHERE id = ?", (int(active), survey_id))

    async def count_submissions(self, survey_id: str) -> int:
        row = await self._fetchone("SELECT COUNT(*) AS c FROM submissions WHERE survey_id = ?", (survey_id,))
        return int(row["c"]) if row else 0

    async def delete_survey(self, survey_id: str) -> bool:
        """Faqat javoblari yo'q so'rovnomani o'chiradi."""
        async with self._write_lock:
            row = await (await self._conn.execute(
                "SELECT COUNT(*) FROM submissions WHERE survey_id = ?", (survey_id,)
            )).fetchone()
            if row and row[0] > 0:
                return False
            cur = await self._conn.execute("DELETE FROM surveys WHERE id = ?", (survey_id,))
            await self._conn.commit()
            return cur.rowcount > 0

    async def surveys_for_grade(self, grade: int, user_id: int) -> tuple[list[Survey], bool]:
        """(hali to'ldirilmagan faol so'rovnomalar, sinf uchun faol so'rovnoma umuman bormi)"""
        active = [s for s in await self.list_surveys(only_active=True) if grade in s.grades]
        rows = await self._fetchall("SELECT survey_id FROM submissions WHERE user_id = ?", (user_id,))
        done = {r["survey_id"] for r in rows}
        return [s for s in active if s.id not in done], bool(active)

    async def has_active_surveys(self) -> bool:
        row = await self._fetchone("SELECT 1 FROM surveys WHERE is_active = 1 LIMIT 1")
        return row is not None

    # ---------- javoblar ----------
    async def save_submission(
        self,
        survey: Survey,
        user_id: int,
        full_name: str,
        grade: str,
        answers: list[str],
    ) -> bool:
        """Barcha javoblarni bitta tranzaksiyada saqlaydi. Takroriy bo'lsa False."""
        async with self._write_lock:
            try:
                cur = await self._conn.execute(
                    "INSERT INTO submissions (survey_id, user_id, full_name, grade) VALUES (?,?,?,?)",
                    (survey.id, user_id, full_name, grade),
                )
                sub_id = cur.lastrowid
                await self._conn.executemany(
                    "INSERT INTO submission_answers (submission_id, q_idx, question, answer) VALUES (?,?,?,?)",
                    [(sub_id, i, survey.questions[i].text, a) for i, a in enumerate(answers)],
                )
                await self._conn.commit()
                return True
            except aiosqlite.IntegrityError:
                await self._conn.rollback()
                return False
            except Exception:
                await self._conn.rollback()
                raise

    # ---------- murojaatlar ----------
    async def link_feedback(self, admin_id: int, admin_msg_id: int, user_id: int, user_name: str) -> None:
        await self._execute(
            "INSERT OR REPLACE INTO feedback_map (admin_id, admin_msg_id, user_id, user_name) VALUES (?,?,?,?)",
            (admin_id, admin_msg_id, user_id, user_name),
        )

    async def get_feedback_target(self, admin_id: int, admin_msg_id: int) -> tuple[int, str] | None:
        row = await self._fetchone(
            "SELECT user_id, user_name FROM feedback_map WHERE admin_id = ? AND admin_msg_id = ?",
            (admin_id, admin_msg_id),
        )
        return (row["user_id"], row["user_name"] or "") if row else None

    # ---------- statistika va eksport ----------
    async def stats(self) -> dict[str, Any]:
        total = (await self._fetchone("SELECT COUNT(*) AS c FROM users"))["c"]
        blocked = (await self._fetchone("SELECT COUNT(*) AS c FROM users WHERE is_blocked = 1"))["c"]
        per_survey = await self._fetchall(
            """SELECT s.id, s.title, s.is_active, COUNT(sub.id) AS c
               FROM surveys s LEFT JOIN submissions sub ON sub.survey_id = s.id
               GROUP BY s.id ORDER BY s.created_at, s.rowid"""
        )
        per_grade = await self._fetchall(
            "SELECT survey_id, grade, COUNT(*) AS c FROM submissions GROUP BY survey_id, grade"
        )
        return {
            "users": total,
            "blocked": blocked,
            "surveys": [dict(r) for r in per_survey],
            "grades": [dict(r) for r in per_grade],
        }

    async def export_data(self) -> list[dict[str, Any]]:
        """Excel uchun: har bir so'rovnoma bo'yicha javob bergan o'quvchilar."""
        result: list[dict[str, Any]] = []
        for survey in await self.list_surveys():
            subs = await self._fetchall(
                """SELECT sub.id, sub.user_id, sub.full_name, sub.grade, sub.created_at,
                          COALESCE(u.username, '') AS username
                   FROM submissions sub LEFT JOIN users u ON u.user_id = sub.user_id
                   WHERE sub.survey_id = ? ORDER BY sub.id""",
                (survey.id,),
            )
            if not subs:
                continue
            ans_rows = await self._fetchall(
                """SELECT a.submission_id, a.q_idx, a.answer
                   FROM submission_answers a JOIN submissions s ON s.id = a.submission_id
                   WHERE s.survey_id = ?""",
                (survey.id,),
            )
            by_sub: dict[int, dict[int, str]] = {}
            for r in ans_rows:
                by_sub.setdefault(r["submission_id"], {})[r["q_idx"]] = r["answer"]
            result.append(
                {
                    "survey": survey,
                    "submissions": [
                        {
                            "user_id": s["user_id"],
                            "username": s["username"],
                            "full_name": s["full_name"],
                            "grade": s["grade"],
                            "created_at": s["created_at"],
                            "answers": [by_sub.get(s["id"], {}).get(i, "") for i in range(len(survey.questions))],
                        }
                        for s in subs
                    ],
                }
            )
        return result


# ====================================================================
# EXCEL_REPORT
# ====================================================================
UZ_TZ = timezone(timedelta(hours=5))  # O'zbekistonda yoz/qish vaqti yo'q

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
ZEBRA_FILL = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
_SIDE = Side(style="thin", color="D9D9D9")
BORDER = Border(left=_SIDE, right=_SIDE, top=_SIDE, bottom=_SIDE)

_BAD_SHEET_CHARS = re.compile(r"[\\/*?:\[\]]")


def _fmt_dt(value: str) -> str:
    try:
        dt = datetime.strptime(value, "%Y-%m-%d %H:%M:%S").replace(tzinfo=timezone.utc)
        return dt.astimezone(UZ_TZ).strftime("%d.%m.%Y %H:%M")
    except (ValueError, TypeError):
        return value or ""


def _set(ws, row: int, col: int, value: Any):
    """Xavfsiz yozish: '=...' formula sifatida talqin qilinmasin, noqonuniy belgilar olib tashlansin."""
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)
    cell = ws.cell(row=row, column=col, value=value)
    if isinstance(value, str):
        cell.data_type = "s"
    return cell


def _sheet_name(survey: Survey, used: set[str]) -> str:
    g = survey.grades
    base = f"{g[0]}-{g[-1]}-sinf" if len(g) > 1 else f"{g[0]}-sinf"
    base = _BAD_SHEET_CHARS.sub("", base)[:28]
    name, n = base, 2
    while name in used:
        name = f"{base} ({n})"
        n += 1
    used.add(name)
    return name


def _style_header(ws, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 45


def build_report(data: list[dict[str, Any]]) -> bytes:
    wb = Workbook()
    summary = wb.active
    summary.title = "Umumiy"
    _set(summary, 1, 1, "So'rovnoma")
    _set(summary, 1, 2, "Sinf")
    _set(summary, 1, 3, "Javoblar soni")
    _style_header(summary, 3)
    summary.row_dimensions[1].height = 25

    used_names = {"Umumiy"}
    srow = 2

    for item in data:
        survey: Survey = item["survey"]
        subs = sorted(item["submissions"], key=lambda s: (grade_key(s["grade"]), s["full_name"].lower()))
        ws = wb.create_sheet(_sheet_name(survey, used_names))

        base_headers = ["T/r", "O'quvchining F.I.SH", "Sinf", "Telegram ID", "Username", "Sana"]
        headers = base_headers + [f"{i}. {q.text}" for i, q in enumerate(survey.questions, start=1)]
        for c, h in enumerate(headers, start=1):
            _set(ws, 1, c, h)
        _style_header(ws, len(headers))

        for r, s in enumerate(subs, start=2):
            username = f"@{s['username']}" if s["username"] else "-"
            values = [r - 1, s["full_name"], s["grade"], s["user_id"], username, _fmt_dt(s["created_at"])]
            values += [a if a else "-" for a in s["answers"]]
            for c, v in enumerate(values, start=1):
                cell = _set(ws, r, c, v)
                cell.border = BORDER
                if r % 2 == 1:
                    cell.fill = ZEBRA_FILL
                center = c in (1, 3, 4, 6)
                cell.alignment = Alignment(
                    horizontal="center" if center else "left", vertical="center", wrap_text=not center
                )

        widths = {1: 6, 2: 30, 3: 8, 4: 14, 5: 18, 6: 17}
        for c in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(c)].width = widths.get(c, 32)
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(subs) + 1}"

        # Umumiy sahifa: sinflar bo'yicha sonlar
        counts: dict[str, int] = {}
        for s in subs:
            counts[s["grade"]] = counts.get(s["grade"], 0) + 1
        first = True
        for grade in sorted(counts, key=grade_key):
            _set(summary, srow, 1, survey.title if first else "")
            _set(summary, srow, 2, grade)
            _set(summary, srow, 3, counts[grade])
            for c in (1, 2, 3):
                summary.cell(row=srow, column=c).border = BORDER
                summary.cell(row=srow, column=c).alignment = Alignment(
                    horizontal="left" if c == 1 else "center", vertical="center", wrap_text=True
                )
            first = False
            srow += 1
        _set(summary, srow, 1, "")
        _set(summary, srow, 2, "Jami")
        _set(summary, srow, 3, len(subs))
        for c in (1, 2, 3):
            cell = summary.cell(row=srow, column=c)
            cell.font = Font(bold=True)
            cell.border = BORDER
            cell.fill = ZEBRA_FILL
            cell.alignment = Alignment(horizontal="center")
        srow += 2

    summary.column_dimensions["A"].width = 60
    summary.column_dimensions["B"].width = 12
    summary.column_dimensions["C"].width = 16

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ====================================================================
# UTILS
# ====================================================================
def esc(text: str | None) -> str:
    """Foydalanuvchi matnini HTML uchun xavfsiz qilish."""
    return escape(text or "", quote=False)


def user_link(user: User) -> str:
    return f'<a href="tg://user?id={user.id}">{esc(user.full_name)}</a>'


def progress_bar(done: int, total: int, width: int = 10) -> str:
    filled = round(width * done / total) if total else 0
    return "▰" * filled + "▱" * (width - filled)


def grade_key(grade: str) -> tuple[int, str]:
    """'10-B' -> (10, 'B') — sinflarni to'g'ri tartiblash uchun."""
    m = re.match(r"(\d+)\s*-?\s*(\w*)", grade)
    return (int(m.group(1)), m.group(2)) if m else (99, grade)


# ====================================================================
# KEYBOARDS
# ====================================================================
class Btn:
    SURVEY = "📝 So'rovnomani to'ldirish"
    SUPPORT = "✉️ Maslahatchiga murojaat"
    CLUBS = "🎭 7 ta rivojlanish klubi"
    COUNCIL = "👥 O'quvchilar kengashi"
    COUNSELOR = "👨‍💼 Maktab maslahatchisi"
    CANCEL = "❌ Bekor qilish"
    # admin (matnlar eski versiya bilan bir xil: eski klaviaturalar ham ishlaydi)
    ADD_SURVEY = "➕ Yangi so'rovnoma qo'shish"
    SURVEYS = "📋 So'rovnomalar"
    STATS = "📈 Statistika"
    EXPORT = "📊 Excel hisobotni yuklash"


def main_keyboard(is_admin: bool = False) -> ReplyKeyboardMarkup:
    rows = [
        [KeyboardButton(text=Btn.SURVEY)],
        [KeyboardButton(text=Btn.SUPPORT)],
        [KeyboardButton(text=Btn.CLUBS), KeyboardButton(text=Btn.COUNCIL)],
        [KeyboardButton(text=Btn.COUNSELOR)],
    ]
    if is_admin:
        rows.append([KeyboardButton(text=Btn.ADD_SURVEY), KeyboardButton(text=Btn.SURVEYS)])
        rows.append([KeyboardButton(text=Btn.STATS), KeyboardButton(text=Btn.EXPORT)])
    return ReplyKeyboardMarkup(keyboard=rows, resize_keyboard=True)


def cancel_keyboard() -> ReplyKeyboardMarkup:
    return ReplyKeyboardMarkup(
        keyboard=[[KeyboardButton(text=Btn.CANCEL)]], resize_keyboard=True
    )


def _cancel_row() -> list[InlineKeyboardButton]:
    return [InlineKeyboardButton(text=Btn.CANCEL, callback_data="cancel_survey")]


def grade_num_keyboard() -> InlineKeyboardMarkup:
    def b(n: int) -> InlineKeyboardButton:
        return InlineKeyboardButton(text=f"{n}-sinf", callback_data=f"num:{n}")

    return InlineKeyboardMarkup(
        inline_keyboard=[[b(5), b(6), b(7)], [b(8), b(9)], [b(10), b(11)], _cancel_row()]
    )


def grade_letter_keyboard(letters: tuple[str, ...]) -> InlineKeyboardMarkup:
    row = [InlineKeyboardButton(text=x, callback_data=f"let:{x}") for x in letters]
    return InlineKeyboardMarkup(inline_keyboard=[row, _cancel_row()])


def survey_pick_keyboard(surveys: list[Survey]) -> InlineKeyboardMarkup:
    rows = [[InlineKeyboardButton(text=s.title[:60], callback_data=f"pick:{s.id}")] for s in surveys]
    rows.append(_cancel_row())
    return InlineKeyboardMarkup(inline_keyboard=rows)


def question_keyboard(q_idx: int, question: Question) -> InlineKeyboardMarkup:
    """callback: ans:<savol_indeksi>:<variant_indeksi> (64 bayt limitidan xavfsiz)"""
    rows = [
        [InlineKeyboardButton(text=opt, callback_data=f"ans:{q_idx}:{i}")]
        for i, opt in enumerate(question.options)
    ]
    rows.append([InlineKeyboardButton(text="❌ So'rovnomani to'xtatish", callback_data="cancel_survey")])
    return InlineKeyboardMarkup(inline_keyboard=rows)


def broadcast_confirm_keyboard() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(
        inline_keyboard=[
            [
                InlineKeyboardButton(text="✅ Yuborish", callback_data="adm_post_ok"),
                InlineKeyboardButton(text="❌ Bekor qilish", callback_data="adm_post_no"),
            ]
        ]
    )


def surveys_admin_keyboard(items: list[tuple[Survey, int]]) -> InlineKeyboardMarkup:
    rows: list[list[InlineKeyboardButton]] = []
    for n, (s, _count) in enumerate(items, start=1):
        rows.append(
            [
                InlineKeyboardButton(
                    text=f"{n}. " + ("⏸ To'xtatish" if s.is_active else "▶️ Yoqish"),
                    callback_data=f"adm_tg:{s.id}",
                ),
                InlineKeyboardButton(text=f"{n}. 🗑 O'chirish", callback_data=f"adm_del:{s.id}"),
            ]
        )
    return InlineKeyboardMarkup(inline_keyboard=rows)


def delete_confirm_keyboard(survey_id: str) -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(
        inline_keyboard=[
            [
                InlineKeyboardButton(text="🗑 Ha, o'chirish", callback_data=f"adm_delok:{survey_id}"),
                InlineKeyboardButton(text="⬅️ Orqaga", callback_data="adm_list"),
            ]
        ]
    )


# ====================================================================
# STATES
# ====================================================================
class SurveyState(StatesGroup):
    waiting_fio = State()
    waiting_grade_num = State()
    waiting_grade_letter = State()
    waiting_survey_pick = State()
    answering = State()


class SupportState(StatesGroup):
    waiting_message = State()


class AdminState(StatesGroup):
    waiting_survey_text = State()
    confirm_broadcast = State()


# ====================================================================
# FILTERS
# ====================================================================
class IsAdmin(BaseFilter):
    async def __call__(self, event: Message | CallbackQuery, settings: Settings) -> bool:
        return event.from_user is not None and event.from_user.id in settings.admin_ids


# ====================================================================
# MIDDLEWARES
# ====================================================================
logger = logging.getLogger(__name__)


class TrackUserMiddleware(BaseMiddleware):
    """Har bir shaxsiy muloqotdagi foydalanuvchini bazada yangilab turadi.

    Baza tozalansa ham foydalanuvchilar e'lon ro'yxatiga qayta tushadi,
    botni blokdan chiqargan foydalanuvchi esa avtomatik tiklanadi.
    """

    def __init__(self, db: Database) -> None:
        self.db = db

    async def __call__(
        self,
        handler: Callable[[TelegramObject, dict[str, Any]], Awaitable[Any]],
        event: TelegramObject,
        data: dict[str, Any],
    ) -> Any:
        user = data.get("event_from_user")
        chat = data.get("event_chat")
        if user and not user.is_bot and (chat is None or chat.type == "private"):
            try:
                await self.db.upsert_user(user.id, user.full_name, user.username)
            except Exception:  # noqa: BLE001 — kuzatuv xatosi asosiy ishni to'xtatmasin
                logger.exception("Foydalanuvchini yangilashda xatolik")
        return await handler(event, data)


# ====================================================================
# MENU
# ====================================================================
menu_router = Router(name="menu")
menu_router.message.filter(F.chat.type == "private")


@menu_router.message(CommandStart())
@menu_router.message(Command("help"))
async def cmd_start(message: Message, state: FSMContext, settings: Settings) -> None:
    await state.clear()
    user = message.from_user
    await message.answer(
        WELCOME.format(name=esc(user.full_name)),
        reply_markup=main_keyboard(user.id in settings.admin_ids),
    )


@menu_router.message(Command("cancel"))
@menu_router.message(F.text == Btn.CANCEL)
async def cancel_any(message: Message, state: FSMContext, settings: Settings) -> None:
    await state.clear()
    await message.answer(
        "Jarayon bekor qilindi va asosiy menyuga qaytdingiz.",
        reply_markup=main_keyboard(message.from_user.id in settings.admin_ids),
    )


@menu_router.callback_query(F.data == "cancel_survey")
async def cancel_survey_cb(cb: CallbackQuery, state: FSMContext, settings: Settings) -> None:
    await state.clear()
    try:
        await cb.message.delete()
    except TelegramBadRequest:
        pass
    await cb.message.answer(
        "So'rovnoma bekor qilindi va asosiy menyuga qaytdingiz.",
        reply_markup=main_keyboard(cb.from_user.id in settings.admin_ids),
    )
    await cb.answer()


@menu_router.message(F.text == Btn.COUNSELOR)
async def show_counselor(message: Message, state: FSMContext) -> None:
    await state.clear()
    await message.answer(ABOUT_COUNSELOR)


@menu_router.message(F.text == Btn.CLUBS)
async def show_clubs(message: Message, state: FSMContext) -> None:
    await state.clear()
    await message.answer(CLUBS_INFO)


@menu_router.message(F.text == Btn.COUNCIL)
async def show_council(message: Message, state: FSMContext) -> None:
    await state.clear()
    await message.answer(COUNCIL_INFO)


@menu_router.message(F.text == Btn.SURVEY)
async def start_survey(message: Message, state: FSMContext, db: Database) -> None:
    await state.clear()
    if not await db.has_active_surveys():
        await message.answer("Hozircha faol so'rovnoma mavjud emas. Keyinroq qayta urinib ko'ring.")
        return
    await message.answer(
        "📝 <b>So'rovnomani to'ldirish uchun avval Ism va Familiyangizni kiriting:</b>\n"
        "<i>(Masalan: Karimov Jasur)</i>\n\n"
        "<i>Jarayonni to'xtatish uchun «❌ Bekor qilish» tugmasini bosing.</i>",
        reply_markup=cancel_keyboard(),
    )
    await state.set_state(SurveyState.waiting_fio)


@menu_router.message(F.text == Btn.SUPPORT)
async def start_support(message: Message, state: FSMContext) -> None:
    await state.clear()
    await message.answer(
        "✍️ <b>Maslahatchiga yo'llamoqchi bo'lgan savol yoki murojaatingizni yozing:</b>\n\n"
        "<i>Matn, rasm, ovozli xabar yoki fayl yuborishingiz mumkin. Xabaringiz maslahatchiga "
        "yetkaziladi va javob shu bot orqali qaytariladi. Bekor qilish uchun pastdagi tugmani bosing.</i>",
        reply_markup=cancel_keyboard(),
    )
    await state.set_state(SupportState.waiting_message)


# ====================================================================
# ADMIN
# ====================================================================
logger = logging.getLogger(__name__)

admin_router = Router(name="admin")
admin_router.message.filter(F.chat.type == "private", IsAdmin())
admin_router.callback_query.filter(IsAdmin())

_background_tasks: set[asyncio.Task] = set()
_COMMAND_PREFIX = re.compile(r"^/\w+(@\w+)?\s*")


def _strip_command(text: str | None) -> str:
    return _COMMAND_PREFIX.sub("", text or "", count=1).strip()


async def _safe_edit(message: Message, text: str, reply_markup=None) -> None:
    try:
        await message.edit_text(text, reply_markup=reply_markup)
    except TelegramBadRequest as e:
        if "not modified" not in str(e).lower():
            raise


# =====================================================================
# 1) Yangi so'rovnoma qo'shish
# =====================================================================
@admin_router.message(F.text == Btn.ADD_SURVEY)
@admin_router.message(Command("add_survey"))
async def add_survey_start(message: Message, state: FSMContext) -> None:
    await state.clear()
    await message.answer(
        "📝 <b>Yangi so'rovnomani bitta xabar shaklida yuboring.</b>\n\n"
        "<b>Namuna:</b>\n"
        f"<pre>{esc(SURVEY_TEMPLATE)}</pre>\n"
        "<i>Qoidalar:</i>\n"
        "• <i>Sinflar: 9-11 yoki 5, 6, 9 ko'rinishida (5–11).</i>\n"
        "• <i>Savoldan keyin «-» bilan variantlar yozilsa — tanlovli savol (kamida 2 ta).</i>\n"
        "• <i>Variant yozilmasa yoki [yozma] bo'lsa — o'quvchi o'zi yozadi.</i>\n\n"
        "Bekor qilish uchun pastdagi tugmani bosing.",
        reply_markup=cancel_keyboard(),
    )
    await state.set_state(AdminState.waiting_survey_text)


# =====================================================================
# 2) So'rovnomalar ro'yxati (yoqish/o'chirish)
# =====================================================================
async def _render_survey_list(db: Database):
    surveys = await db.list_surveys()
    if not surveys:
        return "So'rovnomalar hali yo'q. «➕ Yangi so'rovnoma qo'shish» tugmasidan foydalaning.", None
    items = [(s, await db.count_submissions(s.id)) for s in surveys]
    lines = ["📋 <b>So'rovnomalar</b>", "🟢 faol • 🔴 to'xtatilgan\n"]
    for n, (s, count) in enumerate(items, start=1):
        mark = "🟢" if s.is_active else "🔴"
        grades = ", ".join(map(str, s.grades))
        lines.append(
            f"<b>{n}.</b> {mark} {esc(s.title)}\n"
            f"    Sinflar: {grades} • Savollar: {len(s.questions)} • Javoblar: {count}"
        )
    return "\n".join(lines), surveys_admin_keyboard(items)


@admin_router.message(F.text == Btn.SURVEYS)
@admin_router.message(Command("surveys"))
async def surveys_list(message: Message, state: FSMContext, db: Database) -> None:
    await state.clear()
    text, kb = await _render_survey_list(db)
    await message.answer(text, reply_markup=kb)


@admin_router.callback_query(F.data == "adm_list")
async def cb_list(cb: CallbackQuery, db: Database) -> None:
    text, kb = await _render_survey_list(db)
    await _safe_edit(cb.message, text, kb)
    await cb.answer()


@admin_router.callback_query(F.data.startswith("adm_tg:"))
async def cb_toggle(cb: CallbackQuery, db: Database) -> None:
    survey = await db.get_survey(cb.data.split(":", 1)[1])
    if survey is None:
        await cb.answer("So'rovnoma topilmadi.", show_alert=True)
        return
    await db.set_survey_active(survey.id, not survey.is_active)
    await cb.answer("🟢 Yoqildi" if not survey.is_active else "🔴 To'xtatildi")
    text, kb = await _render_survey_list(db)
    await _safe_edit(cb.message, text, kb)


@admin_router.callback_query(F.data.startswith("adm_del:"))
async def cb_delete_ask(cb: CallbackQuery, db: Database) -> None:
    survey = await db.get_survey(cb.data.split(":", 1)[1])
    if survey is None:
        await cb.answer("So'rovnoma topilmadi.", show_alert=True)
        return
    count = await db.count_submissions(survey.id)
    if count:
        await cb.answer(
            f"Bu so'rovnomada {count} ta javob bor, shuning uchun o'chirib bo'lmaydi. "
            "«To'xtatish» tugmasidan foydalaning.",
            show_alert=True,
        )
        return
    await _safe_edit(
        cb.message,
        f"🗑 <b>{esc(survey.title)}</b> so'rovnomasini o'chirasizmi?",
        delete_confirm_keyboard(survey.id),
    )
    await cb.answer()


@admin_router.callback_query(F.data.startswith("adm_delok:"))
async def cb_delete_ok(cb: CallbackQuery, db: Database) -> None:
    ok = await db.delete_survey(cb.data.split(":", 1)[1])
    await cb.answer("🗑 O'chirildi" if ok else "O'chirib bo'lmadi (javoblar mavjud).", show_alert=not ok)
    text, kb = await _render_survey_list(db)
    await _safe_edit(cb.message, text, kb)


# =====================================================================
# 3) Statistika
# =====================================================================
@admin_router.message(F.text == Btn.STATS)
@admin_router.message(Command("stats"))
async def show_stats(message: Message, state: FSMContext, db: Database) -> None:
    await state.clear()
    st = await db.stats()
    by_survey: dict[str, list[tuple[str, int]]] = {}
    for g in st["grades"]:
        by_survey.setdefault(g["survey_id"], []).append((g["grade"], g["c"]))

    lines = [
        "📈 <b>Statistika</b>",
        f"👤 Foydalanuvchilar: <b>{st['users']}</b> (botni to'xtatganlar: {st['blocked']})\n",
    ]
    for s in st["surveys"]:
        mark = "🟢" if s["is_active"] else "🔴"
        lines.append(f"{mark} <b>{esc(s['title'])}</b>\n    Javoblar: <b>{s['c']}</b>")
        grades = sorted(by_survey.get(s["id"], []), key=lambda x: grade_key(x[0]))
        if grades:
            lines.append("    " + " • ".join(f"{g}: {c}" for g, c in grades))
    text = "\n".join(lines)
    await message.answer(text[:4000])


# =====================================================================
# 4) Excel hisobot
# =====================================================================
@admin_router.message(F.text == Btn.EXPORT)
@admin_router.message(Command("export"))
async def export_excel(message: Message, state: FSMContext, db: Database) -> None:
    await state.clear()
    wait = await message.answer("⏳ So'rovnoma natijalari Excel faylga tayyorlanmoqda...")
    data = await db.export_data()
    if not data:
        await wait.edit_text("Bazada hozircha to'ldirilgan so'rovnomalar mavjud emas.")
        return
    content = await asyncio.to_thread(build_report, data)  # bot bloklanmasligi uchun
    filename = f"Sorovnoma_hisoboti_{datetime.now(UZ_TZ):%Y-%m-%d}.xlsx"
    await message.answer_document(
        BufferedInputFile(content, filename=filename),
        caption="📊 <b>Maktab so'rovnomalari bo'yicha Excel hisobot</b>\n"
        "<i>Har bir so'rovnoma alohida varaqda, 1-varaq — umumiy jadval.</i>",
    )
    await wait.delete()


# =====================================================================
# 5) E'lon yuborish (/post) — oldindan ko'rish va tasdiqlash bilan
# =====================================================================
def _announcement(text: str) -> str:
    return f"📢 <b>Maktab maslahatchisidan umumiy e'lon:</b>\n━━━━━━━━━━━━━━━━━━━━\n{text}"


@admin_router.message(Command("post"))
async def post_command(message: Message, state: FSMContext, db: Database) -> None:
    body = _strip_command(message.html_text)  # Telegram'dagi qalin/kursiv formatlash saqlanadi
    if not body:
        await message.answer(
            "⚠️ <b>Matn kiritilmadi!</b>\nFoydalanish:\n"
            "<code>/post Barcha o'quvchilar diqqatiga! Ertaga tadbir bo'ladi.</code>"
        )
        return
    text = _announcement(body)
    try:
        await message.answer(text)  # ko'rinishi
    except TelegramBadRequest:
        await message.answer("⚠️ Matnni yuborib bo'lmadi (formatlashda xato bor). Soddaroq yozib ko'ring.")
        return
    recipients = len(await db.get_broadcast_ids())
    await state.set_state(AdminState.confirm_broadcast)
    await state.update_data(broadcast_text=text)
    await message.answer(
        f"👆 Shu e'lon <b>{recipients} ta</b> foydalanuvchiga yuboriladi. Tasdiqlaysizmi?",
        reply_markup=broadcast_confirm_keyboard(),
    )


async def _send_one(bot: Bot, user_id: int, text: str, db: Database) -> str:
    for _ in range(3):
        try:
            await bot.send_message(user_id, text)
            return "ok"
        except TelegramRetryAfter as e:
            await asyncio.sleep(e.retry_after + 1)
        except TelegramForbiddenError:
            await db.mark_blocked(user_id)
            return "blocked"
        except TelegramBadRequest as e:
            if "chat not found" in str(e).lower():
                await db.mark_blocked(user_id)
                return "blocked"
            return "failed"
        except TelegramAPIError:
            return "failed"
    return "failed"


async def _run_broadcast(bot: Bot, db: Database, settings: Settings, text: str, admin_id: int) -> None:
    ids = await db.get_broadcast_ids()
    stats = {"ok": 0, "blocked": 0, "failed": 0}
    for uid in ids:
        stats[await _send_one(bot, uid, text, db)] += 1
        await asyncio.sleep(settings.broadcast_delay)
    await bot.send_message(
        admin_id,
        "✅ <b>E'lon tarqatildi!</b>\n\n"
        f"• Yetkazildi: <b>{stats['ok']}</b>\n"
        f"• Botni to'xtatganlar: <b>{stats['blocked']}</b>\n"
        f"• Boshqa xatolar: <b>{stats['failed']}</b>",
    )


@admin_router.callback_query(AdminState.confirm_broadcast, F.data == "adm_post_ok")
async def broadcast_confirm(cb: CallbackQuery, state: FSMContext, bot: Bot, db: Database, settings: Settings) -> None:
    data = await state.get_data()
    text = data.get("broadcast_text")
    await state.clear()
    if not text:
        await cb.answer("E'lon matni topilmadi.", show_alert=True)
        return
    await _safe_edit(cb.message, "⏳ E'lon yuborilmoqda... Tugagach hisobot keladi.")
    await cb.answer()
    task = asyncio.create_task(_run_broadcast(bot, db, settings, text, cb.from_user.id))
    _background_tasks.add(task)  # GC o'chirib yubormasligi uchun
    task.add_done_callback(_background_tasks.discard)


@admin_router.callback_query(F.data == "adm_post_no")
async def broadcast_cancel(cb: CallbackQuery, state: FSMContext) -> None:
    await state.clear()
    await _safe_edit(cb.message, "❌ E'lon yuborish bekor qilindi.")
    await cb.answer()


# =====================================================================
# 6) Bitta foydalanuvchiga xabar (/sms)
# =====================================================================
@admin_router.message(Command("sms"))
async def sms_command(message: Message, bot: Bot, db: Database) -> None:
    parts = _strip_command(message.html_text).split(maxsplit=1)
    if len(parts) < 2 or not parts[0].lstrip("-").isdigit():
        await message.answer(
            "⚠️ <b>Noto'g'ri format!</b>\nFoydalanish:\n<code>/sms 123456789 Xonamga kirib o'tgin</code>"
        )
        return
    target_id, text = int(parts[0]), parts[1]
    try:
        await bot.send_message(
            target_id,
            f"👨‍💼 <b>Maktab maslahatchisidan shaxsiy xabar:</b>\n━━━━━━━━━━━━━━━━━━━━\n{text}",
        )
        await message.answer(f"✅ Xabar <code>{target_id}</code> ga yetkazildi.")
    except TelegramForbiddenError:
        await db.mark_blocked(target_id)
        await message.answer("❌ Foydalanuvchi botni to'xtatgan.")
    except TelegramAPIError as e:
        await message.answer(f"❌ Xabarni yuborib bo'lmadi: {esc(str(e))}")


# =====================================================================
# 7) Murojaatga Reply orqali javob (matn, rasm, ovoz, fayl — hammasi)
# =====================================================================
@admin_router.message(F.reply_to_message)
async def reply_to_user(message: Message, bot: Bot, db: Database) -> None:
    target = await db.get_feedback_target(message.from_user.id, message.reply_to_message.message_id)
    if target is None:
        raise SkipHandler  # bu murojaat emas — boshqa handlerlarga o'tsin
    user_id, name = target
    header = "👨‍💼 <b>Maktab maslahatchisidan javob:</b>\n━━━━━━━━━━━━━━━━━━━━"
    try:
        if message.text:
            await bot.send_message(user_id, f"{header}\n{message.html_text}")
        else:
            await bot.send_message(user_id, header)
            await message.copy_to(user_id)
        await message.reply(f"✅ Javobingiz <b>{esc(name)}</b>ga yuborildi!")
    except TelegramForbiddenError:
        await db.mark_blocked(user_id)
        await message.reply("❌ Javobni yuborib bo'lmadi: foydalanuvchi botni to'xtatgan.")
    except TelegramAPIError as e:
        await message.reply(f"❌ Javobni yuborib bo'lmadi: {esc(str(e))}")


# =====================================================================
# 8) So'rovnoma matnini qabul qilish (holat handlerlari oxirida)
# =====================================================================
@admin_router.message(AdminState.waiting_survey_text, F.text)
async def save_survey(message: Message, state: FSMContext, db: Database, settings: Settings) -> None:
    title, grades, questions, errors = parse_survey_text(message.text)
    if errors:
        shown = "\n".join(f"• {esc(e)}" for e in errors[:8])
        await message.answer(f"⚠️ <b>Matnda xatolar bor:</b>\n{shown}\n\nTuzatib qayta yuboring yoki bekor qiling.")
        return

    survey_id = f"s_{uuid.uuid4().hex[:8]}"
    await db.add_survey(survey_id, title, grades, questions)
    await state.clear()
    await message.answer(
        "✅ <b>Yangi so'rovnoma saqlandi!</b>\n\n"
        f"📌 <b>Nomi:</b> {esc(title)}\n"
        f"👥 <b>Sinflar:</b> {', '.join(map(str, grades))}\n"
        f"❓ <b>Savollar soni:</b> {len(questions)} ta\n\n"
        "O'quvchilar so'rovnomani to'ldirishda sinfiga mos so'rovnomani ko'radi. "
        "Bir sinfga bir nechta so'rovnoma tegishli bo'lsa, tanlash taklif etiladi.",
        reply_markup=main_keyboard(True),
    )


@admin_router.message(AdminState.waiting_survey_text)
async def survey_text_expected(message: Message) -> None:
    await message.answer("⚠️ Iltimos, so'rovnomani matn ko'rinishida yuboring.")


# ====================================================================
# SURVEY
# ====================================================================
logger = logging.getLogger(__name__)

survey_router = Router(name="survey")
survey_router.message.filter(F.chat.type == "private")

MAX_TEXT_ANSWER = 1000


# ---------------- yordamchi funksiyalar ----------------
def _valid_fio(fio: str) -> bool:
    return 5 <= len(fio) <= 80 and len(fio.split()) >= 2 and not any(ch.isdigit() for ch in fio)


async def _send_question(bot: Bot, chat_id: int, survey: Survey, idx: int) -> None:
    q = survey.questions[idx]
    total = len(survey.questions)
    text = (
        f"📊 <b>Savol {idx + 1}/{total}</b>  {progress_bar(idx, total)}\n"
        f"━━━━━━━━━━━━━━━━━━━━\n<b>{esc(q.text)}</b>"
    )
    if q.type == "choice":
        await bot.send_message(chat_id, text, reply_markup=question_keyboard(idx, q))
    else:
        await bot.send_message(chat_id, text + "\n\n✍️ <i>Javobingizni xabar sifatida yozib yuboring:</i>")


async def _begin(bot: Bot, chat_id: int, state: FSMContext, survey: Survey) -> None:
    await state.update_data(survey_id=survey.id, q_idx=0, answers=[])
    await state.set_state(SurveyState.answering)
    await _send_question(bot, chat_id, survey, 0)


async def _advance(bot: Bot, db: Database, settings: Settings, state: FSMContext, user_id: int, answer: str) -> None:
    data = await state.get_data()
    survey = await db.get_survey(data["survey_id"])
    keyboard = main_keyboard(user_id in settings.admin_ids)
    if survey is None:
        await state.clear()
        await bot.send_message(user_id, "Bu so'rovnoma o'chirib yuborilgan.", reply_markup=keyboard)
        return

    answers = [*data["answers"], answer]
    if len(answers) < len(survey.questions):
        await state.update_data(q_idx=len(answers), answers=answers)
        await _send_question(bot, user_id, survey, len(answers))
        return

    fio = data["fio"]
    saved = await db.save_submission(survey, user_id, fio, data["full_grade"], answers)
    await state.clear()
    if saved:
        await bot.send_message(
            user_id,
            f"✅ <b>Katta rahmat, {esc(fio)}!</b>\n\n"
            "Siz so'rovnomani to'liq yakunladingiz. Barcha javoblaringiz maslahatchiga yetkazildi.",
            reply_markup=keyboard,
        )
    else:
        await bot.send_message(user_id, "Siz bu so'rovnomani allaqachon to'ldirgansiz.", reply_markup=keyboard)


# ---------------- F.I.SH ----------------
@survey_router.message(SurveyState.waiting_fio, F.text)
async def process_fio(message: Message, state: FSMContext) -> None:
    fio = " ".join(message.text.split())
    if not _valid_fio(fio):
        await message.answer(
            "⚠️ Iltimos, ism va familiyangizni to'liq (raqamsiz) yozing.\n<i>Masalan: Karimov Jasur</i>"
        )
        return
    await state.update_data(fio=fio)
    await state.set_state(SurveyState.waiting_grade_num)
    await message.answer(f"Rahmat, <b>{esc(fio)}</b>!\nEndi sinfingizni tanlang:", reply_markup=grade_num_keyboard())


# ---------------- sinf ----------------
@survey_router.callback_query(SurveyState.waiting_grade_num, F.data.startswith("num:"))
async def process_grade_num(cb: CallbackQuery, state: FSMContext, settings: Settings) -> None:
    grade_num = int(cb.data.split(":")[1])
    if not 5 <= grade_num <= 11:
        await cb.answer()
        return
    await state.update_data(grade_num=grade_num)
    await state.set_state(SurveyState.waiting_grade_letter)
    await cb.message.edit_text(
        f"Tanlandi: <b>{grade_num}-sinf</b>\nSinfingiz harfini tanlang:",
        reply_markup=grade_letter_keyboard(settings.grade_letters),
    )
    await cb.answer()


@survey_router.callback_query(SurveyState.waiting_grade_letter, F.data.startswith("let:"))
async def process_grade_letter(
    cb: CallbackQuery, state: FSMContext, bot: Bot, db: Database, settings: Settings
) -> None:
    letter = cb.data.split(":", 1)[1]
    data = await state.get_data()
    grade_num = data["grade_num"]
    full_grade = f"{grade_num}-{letter}"
    await state.update_data(full_grade=full_grade)

    available, any_active = await db.surveys_for_grade(grade_num, cb.from_user.id)
    keyboard = main_keyboard(cb.from_user.id in settings.admin_ids)

    if not available:
        try:
            await cb.message.delete()
        except TelegramBadRequest:
            pass
        text = (
            "Siz bu sinf uchun mavjud barcha so'rovnomalarni allaqachon to'ldirgansiz. Rahmat! 🙏"
            if any_active
            else f"{grade_num}-sinf uchun hozircha faol so'rovnoma mavjud emas."
        )
        await state.clear()
        await bot.send_message(cb.from_user.id, text, reply_markup=keyboard)
        await cb.answer()
        return

    try:
        await cb.message.delete()
    except TelegramBadRequest:
        pass

    if len(available) == 1:
        await _begin(bot, cb.from_user.id, state, available[0])
    else:
        await state.set_state(SurveyState.waiting_survey_pick)
        await bot.send_message(
            cb.from_user.id,
            f"<b>{esc(full_grade)}</b> sinfi uchun bir nechta so'rovnoma mavjud. Birini tanlang:",
            reply_markup=survey_pick_keyboard(available),
        )
    await cb.answer()


@survey_router.callback_query(SurveyState.waiting_survey_pick, F.data.startswith("pick:"))
async def process_pick(cb: CallbackQuery, state: FSMContext, bot: Bot, db: Database) -> None:
    survey = await db.get_survey(cb.data.split(":", 1)[1])
    if survey is None or not survey.is_active:
        await cb.answer("Bu so'rovnoma endi mavjud emas.", show_alert=True)
        return
    try:
        await cb.message.delete()
    except TelegramBadRequest:
        pass
    await _begin(bot, cb.from_user.id, state, survey)
    await cb.answer()


# ---------------- savollarga javob ----------------
@survey_router.callback_query(SurveyState.answering, F.data.startswith("ans:"))
async def handle_choice(cb: CallbackQuery, state: FSMContext, bot: Bot, db: Database, settings: Settings) -> None:
    try:
        _, q_s, o_s = cb.data.split(":")
        q_pos, opt_pos = int(q_s), int(o_s)
    except ValueError:
        await cb.answer()
        return

    data = await state.get_data()
    survey = await db.get_survey(data["survey_id"])
    if survey is None:
        await state.clear()
        await cb.answer("So'rovnoma topilmadi.", show_alert=True)
        return
    if q_pos != data["q_idx"]:
        await cb.answer("Bu savolga allaqachon javob berilgan.")
        return

    question = survey.questions[q_pos]
    if question.type != "choice" or not 0 <= opt_pos < len(question.options):
        await cb.answer()
        return

    answer = question.options[opt_pos]
    try:  # tanlangan javobni xabarda qoldirib, tugmalarni olib tashlaymiz
        await cb.message.edit_text(f"<b>{esc(question.text)}</b>\n\n✅ {esc(answer)}", reply_markup=None)
    except TelegramBadRequest:
        pass
    await cb.answer()
    await _advance(bot, db, settings, state, cb.from_user.id, answer)


@survey_router.message(SurveyState.answering, F.text)
async def handle_text(message: Message, state: FSMContext, bot: Bot, db: Database, settings: Settings) -> None:
    data = await state.get_data()
    survey = await db.get_survey(data["survey_id"])
    if survey is None:
        await state.clear()
        await message.answer("Bu so'rovnoma o'chirib yuborilgan.", reply_markup=main_keyboard(message.from_user.id in settings.admin_ids))
        return

    question = survey.questions[data["q_idx"]]
    if question.type == "choice":
        await message.answer("👆 Iltimos, yuqoridagi tugmalardan birini tanlang.")
        return

    answer = message.text.strip()
    if len(answer) > MAX_TEXT_ANSWER:
        await message.answer(f"⚠️ Javob juda uzun. Iltimos, {MAX_TEXT_ANSWER} belgidan qisqaroq yozing.")
        return
    await _advance(bot, db, settings, state, message.from_user.id, answer)


# ---------------- noto'g'ri turdagi kiritish ----------------
@survey_router.message(StateFilter(SurveyState.waiting_fio, SurveyState.answering))
async def expect_text(message: Message) -> None:
    await message.answer("⚠️ Iltimos, javobni oddiy matn ko'rinishida yuboring.")


@survey_router.message(
    StateFilter(SurveyState.waiting_grade_num, SurveyState.waiting_grade_letter, SurveyState.waiting_survey_pick)
)
async def expect_button(message: Message) -> None:
    await message.answer("👆 Iltimos, yuqoridagi tugmalardan birini tanlang.")


# ====================================================================
# SUPPORT
# ====================================================================
logger = logging.getLogger(__name__)

support_router = Router(name="support")
support_router.message.filter(F.chat.type == "private")

THROTTLE_SECONDS = 15
_last_sent: dict[int, float] = {}


@support_router.message(SupportState.waiting_message)
async def process_inquiry(message: Message, state: FSMContext, bot: Bot, db: Database, settings: Settings) -> None:
    user = message.from_user
    now = time.monotonic()
    if now - _last_sent.get(user.id, 0.0) < THROTTLE_SECONDS:
        await message.answer("⏳ Iltimos, keyingi murojaatni yuborishdan oldin biroz kuting.")
        return

    username = f" @{esc(user.username)}" if user.username else ""
    header = (
        "📬 <b>Yangi murojaat</b>\n"
        f"👤 {user_link(user)}{username}\n"
        f"🆔 <code>{user.id}</code>\n"
        "━━━━━━━━━━━━━━━━━━━━\n"
        "<i>Javob berish uchun shu xabarga (yoki quyidagi nusxasiga) Reply qiling.</i>"
    )

    delivered = 0
    for admin_id in settings.admin_ids:
        try:
            head = await bot.send_message(admin_id, header)
            copy = await message.copy_to(admin_id)
            await db.link_feedback(admin_id, head.message_id, user.id, user.full_name)
            await db.link_feedback(admin_id, copy.message_id, user.id, user.full_name)
            delivered += 1
        except TelegramAPIError:
            logger.warning("Murojaatni adminga (%s) yetkazib bo'lmadi", admin_id, exc_info=True)

    if not delivered:
        await message.answer("Xabarni yuborishda xatolik yuz berdi. Keyinroq urinib ko'ring.")
        return

    _last_sent[user.id] = now
    await state.clear()
    await message.answer(
        "✅ <b>Xabaringiz maktab maslahatchisiga yetkazildi!</b>\nTez orada bot orqali sizga javob qaytariladi.",
        reply_markup=main_keyboard(user.id in settings.admin_ids),
    )


# ====================================================================
# FALLBACK
# ====================================================================
fallback_router = Router(name="fallback")


@fallback_router.message(F.chat.type == "private", StateFilter(None))
async def unknown_message(message: Message, settings: Settings) -> None:
    await message.answer(
        "Quyidagi menyudan kerakli bo'limni tanlang 👇",
        reply_markup=main_keyboard(message.from_user.id in settings.admin_ids),
    )


@fallback_router.callback_query()
async def stale_callback(cb: CallbackQuery) -> None:
    await cb.answer("Bu tugma eskirgan. /start buyrug'ini bosing.", show_alert=True)


# ====================================================================
# WEB
# ====================================================================
async def _ping(_: web.Request) -> web.Response:
    return web.Response(text="OK")


async def start_web_server(port: int) -> web.AppRunner:
    app = web.Application()
    app.router.add_get("/", _ping)
    app.router.add_get("/health", _ping)
    runner = web.AppRunner(app)
    await runner.setup()
    await web.TCPSite(runner, "0.0.0.0", port).start()
    return runner


# ====================================================================
# MAIN
# ====================================================================
logger = logging.getLogger(__name__)

USER_COMMANDS = [
    BotCommand(command="start", description="Asosiy menyu"),
    BotCommand(command="cancel", description="Jarayonni bekor qilish"),
]
ADMIN_COMMANDS = USER_COMMANDS + [
    BotCommand(command="stats", description="Statistika"),
    BotCommand(command="surveys", description="So'rovnomalarni boshqarish"),
    BotCommand(command="add_survey", description="Yangi so'rovnoma qo'shish"),
    BotCommand(command="export", description="Excel hisobot"),
    BotCommand(command="post", description="Hammaga e'lon yuborish"),
    BotCommand(command="sms", description="Foydalanuvchiga xabar"),
]


async def on_error(event: ErrorEvent, bot: Bot) -> bool:
    logger.exception("Ishlov berilmagan xatolik: %s", event.exception, exc_info=event.exception)
    update = event.update
    chat_id = None
    if update.message:
        chat_id = update.message.chat.id
    elif update.callback_query:
        chat_id = update.callback_query.from_user.id
        try:
            await update.callback_query.answer()
        except Exception:  # noqa: BLE001
            pass
    if chat_id:
        try:
            await bot.send_message(chat_id, "⚠️ Kutilmagan xatolik yuz berdi. Iltimos, /start ni bosib qayta urinib ko'ring.")
        except Exception:  # noqa: BLE001
            pass
    return True


async def main() -> None:
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s | %(levelname)-7s | %(name)s | %(message)s",
    )
    settings = load_settings()
    db = await Database.connect(settings.db_path)

    bot = Bot(token=settings.bot_token, default=DefaultBotProperties(parse_mode=ParseMode.HTML))
    dp = Dispatcher(
        storage=MemoryStorage(),
        events_isolation=SimpleEventIsolation(),  # bir foydalanuvchining hodisalari ketma-ket ishlanadi
        db=db,
        settings=settings,
    )

    tracker = TrackUserMiddleware(db)
    dp.message.outer_middleware(tracker)
    dp.callback_query.outer_middleware(tracker)

    # Tartib muhim: menyu -> admin -> holat handlerlari -> fallback
    dp.include_routers(menu_router, admin_router, survey_router, support_router, fallback_router)
    dp.errors.register(on_error)

    runner = await start_web_server(settings.port)
    try:
        await bot.set_my_commands(USER_COMMANDS)
        for admin_id in settings.admin_ids:
            try:
                await bot.set_my_commands(ADMIN_COMMANDS, scope=BotCommandScopeChat(chat_id=admin_id))
            except Exception:  # noqa: BLE001 — admin hali botni ishga tushirmagan bo'lishi mumkin
                logger.warning("Admin (%s) uchun buyruqlar o'rnatilmadi", admin_id)

        await bot.delete_webhook(drop_pending_updates=False)
        logger.info("Bot ishga tushdi.")
        await dp.start_polling(bot, allowed_updates=dp.resolve_used_update_types())
    finally:
        await runner.cleanup()
        await bot.session.close()
        await db.close()


def run() -> None:
    try:
        asyncio.run(main())
    except (KeyboardInterrupt, SystemExit):
        logger.info("Bot to'xtatildi.")


if __name__ == "__main__":
    run()
