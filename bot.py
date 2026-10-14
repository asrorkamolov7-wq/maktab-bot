import asyncio
import os
import aiosqlite
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from aiohttp import web
from aiogram import Bot, Dispatcher, F
from aiogram.enums import ParseMode
from aiogram.client.default import DefaultBotProperties
from aiogram.filters import CommandStart, Command
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.types import (
    Message,
    CallbackQuery,
    ReplyKeyboardMarkup,
    KeyboardButton,
    ReplyKeyboardRemove,
    InlineKeyboardMarkup,
    InlineKeyboardButton,
    FSInputFile
)

# ================= SOZLAMALAR =================
BOT_TOKEN = "8619703357:AAHIGSxKdB2k2Lw643-WzsIOQf_OrGoUgTE"
ADMIN_ID = 7680698926
DB_NAME = "school_bot.db"

# ================= MAKTAB MASLAHATCHISI MA'LUMOTLARI =================
ABOUT_COUNSELOR = """
👨‍💼 <b>Maktab maslahatchisi kabineti</b>
━━━━━━━━━━━━━━━━━━━━
• <b>F.I.SH:</b> Kamolov Asror
• <b>Lavozimi:</b> Maktab maslahatchisi
• <b>Xona:</b> 2-qavat, Maslahatchi xonasi
• <b>Qabul vaqti:</b> 09:00 dan 16:00 gacha
• <b>Bog'lanish uchun telefon:</b> +998 33 443 55 77

<i>Har qanday shaxsiy masala, kasb tanlash yoki qiyinchiliklar bo'yicha bemalol murojaat qilishingiz mumkin. Barcha suhbatlar to'liq sir saqlanadi.</i>
"""

# ================= 7 TA KLUB =================
CLUBS_INFO = """
🎭 <b>Maktabda faoliyat yuritayotgan 7 ta rivojlanish klubi:</b>
━━━━━━━━━━━━━━━━━━━━
1. 🎭 <b>Turon teatr</b> — Sahna madaniyati, aktyorlik mahorati va ijodiy qobiliyatlarni rivojlantirish.
2. 🎵 <b>Iqtidor ansambli</b> — Musiqa, milliy va zamonaviy ijrochilik sanʼati.
3. 📜 <b>Jadidlar izidan</b> — Tarixiy meros, kitobxonlik va maʼrifatparvarlik tadbirlari.
4. 🌿 <b>Eco-schools Uzbekistan</b> — Ekologik savodxonlik, tabiatni asrash va atrof-muhit loyihalari.
5. 🌍 <b>Xorijiy tillar</b> — Chet tillarida erkin muloqot va soʻzlashuv klubi.
6. 🎤 <b>Debat</b> — Notiqlik, tanqidiy fikrlash va munozara madaniyati.
7. 💻 <b>Raqamli avlod qizlari</b> — Qizlar uchun IT, zamonaviy kasblar va media savodxonlik.

<i>Klublarga a'zo bo'lish uchun maktab maslahatchisi xonasiga (2-qavat) murojaat qiling.</i>
"""

# ================= O'QUVCHILAR KENGASHI =================
COUNCIL_INFO = """
👥 <b>O'quvchilar kengashi tashkiliy tuzilmasi</b>
━━━━━━━━━━━━━━━━━━━━
👑 <b>Kelajak lideri</b>
O'quvchilar kengashi faoliyatiga umumiy rahbarlik qiladi.

📌 <b>Asosiy yo'nalish koordinatorlari va vazifalari:</b>

🎓 <b>Ta'lim koordinatori</b>
Ta'lim sifatini oshirish, fan va til to'garaklarini tashkil etish, o'quvchilarning o'zlashtirishiga yordam berish.

🎭 <b>Madaniyat koordinatori</b>
Maktabning madaniy tadbirlarini tashkil etish, milliy qadriyatlarni targ'ib qilish, "Turon maktab teatri" va "Iqtidor ansambli" faoliyatini qo'llab-quvvatlash.

💻 <b>IT va sun'iy intellekt koordinatori</b>
Raqamli texnologiyalarni targ'ib qilish, IT ko'nikmalarini rivojlantirish, innovatsion loyihalar va tanlovlarni tashkil etish.

🏃‍♂️ <b>Sport koordinatori</b>
Sport tadbirlarini tashkil etish, sog'lom turmush tarzini targ'ib qilish, sport klublari va tanlovlarni rivojlantirish.

🌿 <b>Ekologiya koordinatori</b>
Ekologik loyihalar va aksiyalarni tashkil etish, tabiatni muhofaza qilish, ekologik madaniyatni targ'ib qilish.

📢 <b>Media koordinatori</b>
Maktab hayotini ommaviy axborotda yoritish, ijtimoiy tarmoqlarni yuritish, media loyihalarini tashkil etish.

👥 <b>Ijtimoiy tashabbuslar koordinatori</b>
O'quvchilar tashabbuslarini qo'llab-quvvatlash, ijtimoiy loyihalar va volontyorlik faoliyatini tashkil etish, inklyuziv muhitni rivojlantirish.

━━━━━━━━━━━━━━━━━━━━
👥 <b>5–11-sinf vakillari:</b> Sinf o'quvchilari tomonidan saylanadi.
☘️ <b>Klublar:</b> Maktabda faoliyat yurituvchi tarmoq to'garaklari.
"""

# ================= SO'ROVNOMALAR RO'YXATI =================
SURVEYS = [
    {
        "id": "survey_5_8",
        "title": "«Mening qiziqishlarim va kelajakdagi kasbim» (5–8-sinf)",
        "grades": [5, 6, 7, 8],
        "questions": [
            {
                "type": "choice",
                "question": "1. Maktabda qaysi fanlarni o'rganish sizga ko'proq yoqadi?",
                "options": ["Matematika", "Ona tili va adabiyot", "Tarix", "Chet tili", "Informatika", "Biologiya", "Fizika", "Jismoniy tarbiya", "Musiqa / Tasviriy san'at"]
            },
            {
                "type": "choice",
                "question": "2. Bo'sh vaqtingizda nimalar bilan shug'ullanishni yoqtirasiz?",
                "options": ["Kitob o'qish", "Sport", "Rasm chizish", "Kompyuter va IT", "Video / montaj / dizayn", "Qo'shiq / cholg'u", "Pazandachilik / hunar"]
            },
            {
                "type": "choice",
                "question": "3. O'zingizda qaysi qobiliyat kuchliroq deb o'ylaysiz?",
                "options": ["Hisoblash va mantiq", "Chiroyli gapirish", "Texnika va kompyuter", "Musiqa va ijod", "Sport", "Tashkilotchilik", "Tinglash va yordam berish", "Hali bilmayman"]
            },
            {
                "type": "choice",
                "question": "4. Yangi narsani o'rganishda qaysi usul sizga ko'proq yoqadi?",
                "options": ["O'qib o'rganish", "Amalda bajarib ko'rish", "Video orqali ko'rish", "Ustozdan so'rab", "Do'stlar bilan birgalikda", "Tajriba qilib ko'rish"]
            },
            {
                "type": "choice",
                "question": "5. Quyidagi faoliyatlardan qaysi biri sizga eng qiziq?",
                "options": ["Kompyuter / IT", "Tibbiyot", "Texnika va muhandislik", "San'at va ijod", "Sport", "Ta'lim va pedagogika", "Biznes va tadbirkorlik", "Huquq / Harbiy"]
            },
            {
                "type": "text",
                "question": "6. Kelajakda qaysi kasb yoki kasblarga qiziqasiz? Nomlarini yozing:"
            },
            {
                "type": "choice",
                "question": "7. Bu kasbga qiziqishingizga nima sabab bo'lgan?",
                "options": ["Bu kasb menga yoqadi", "Tanishim shu sohada ishlaydi", "Bolalikdagi orzuim", "Daromadi yaxshi", "Odamlarga yordam berish", "Hali aniq bilmayman"]
            },
            {
                "type": "choice",
                "question": "8. Siz ko'proq qanday ishni bajarishni xohlaysiz?",
                "options": ["Odamlar bilan ishlash", "Kompyuter bilan ishlash", "Texnika va qurilmalar", "Hayvonlar va tabiat", "Ijodiy ishlar", "Jamoa bilan ishlash", "Mustaqil ishlash"]
            },
            {
                "type": "choice",
                "question": "9. Qaysi to'garak yoki klub sizni ko'proq qiziqtiradi?",
                "options": ["Turon teatr", "Iqtidor ansambli", "Jadidlar izidan", "Eco-schools Uzbekistan", "Xorijiy tillar", "Debat klubi", "Raqamli avlod qizlari"]
            },
            {
                "type": "text",
                "question": "10. Maktabda yangi to'garak ochilsa, qaysi yo'nalishda bo'lishini xohlardingiz?"
            },
            {
                "type": "choice",
                "question": "11. O'zingizni kelajakda qanday inson sifatida tasavvur qilasiz?",
                "options": ["Yaxshi mutaxassis", "Tadbirkor", "Olim / Tadqiqotchi", "Ijodkor", "Sportchi", "Rahbar", "O'qituvchi"]
            },
            {
                "type": "choice",
                "question": "12. Sizga kasb tanlashda kimning maslahati ko'proq ta'sir qiladi?",
                "options": ["Ota-onam", "O'qituvchim", "Sinfdoshlarim / do'stlarim", "Maktab maslahatchisi", "O'zim mustaqil tanlayman"]
            },
            {
                "type": "text",
                "question": "13. Kelajakdagi qaysi kasb haqida ko'proq ma'lumot olishni xohlaysiz?"
            },
            {
                "type": "text",
                "question": "14. Siz maktab hayotini yaxshilash uchun qanday yangi taklif bera olasiz?"
            },
            {
                "type": "text",
                "question": "15. O'zingiz haqingizda biz bilishimiz kerak bo'lgan boshqa qiziqishingiz bormi?"
            }
        ]
    },
    {
        "id": "survey_9_11",
        "title": "«Kasb tanlash va oliygohga tayyorgarlik» (9–11-sinf)",
        "grades": [9, 10, 11],
        "questions": [
            {
                "type": "choice",
                "question": "1. Kelajakda qaysi kasb egasi bo'lishni xohlaysiz?",
                "options": ["Tibbiyot", "Muhandislik / Arxitektura", "IT sohasi", "Iqtisodiyot / Biznes", "Ta'lim / Pedagogika", "Huquqshunoslik", "Harbiy soha", "Boshqa"]
            },
            {
                "type": "choice",
                "question": "2. Sizni qaysi fanlar ko'proq qiziqtiradi?",
                "options": ["Matematika", "Fizika", "Kimyo / Biologiya", "Ingliz tili / Chet tillari", "Tarix va huquq", "Ona tili va adabiyot"]
            },
            {
                "type": "choice",
                "question": "3. Maktabni tugatgach o'qishni davom ettirishni rejalashtiryapsizmi?",
                "options": ["Ha, albatta", "Yo'q", "Hali aniq emas"]
            },
            {
                "type": "choice",
                "question": "4. Qayerda o'qishni xohlaysiz?",
                "options": ["O'zbekiston oliy ta'lim muassasalarida", "Xorijiy davlat universitetlarida", "Ikkalasini ham ko'rib chiqyapman"]
            },
            {
                "type": "choice",
                "question": "5. Qaysi davlat universitetlarida o'qish sizni qiziqtiradi?",
                "options": ["O'zbekiston", "AQSh", "Buyuk Britaniya", "Koreya", "Turkiya", "Yevropa davlatlari", "Rossiya"]
            },
            {
                "type": "choice",
                "question": "6. Xorijiy universitetlarga hujjat topshirishni rejalashtiryapsizmi?",
                "options": ["Ha", "Yo'q", "Hali o'ylab ko'ryapman"]
            },
            {
                "type": "choice",
                "question": "7. Qaysi xalqaro test yoki sertifikatlarga tayyorlanyapsiz?",
                "options": ["IELTS", "SAT", "CEFR", "TOEFL", "Hali boshlamaganman"]
            },
            {
                "type": "choice",
                "question": "8. Ingliz tilini bilish darajangiz qanday?",
                "options": ["Boshlang'ich", "O'rta", "Yaxshi (B2 daraja)", "Juda yaxshi (C1 daraja)"]
            },
            {
                "type": "text",
                "question": "9. Hozirda qaysi to'garak yoki o'quv markazlariga qatnashyapsiz?"
            },
            {
                "type": "choice",
                "question": "10. Siz qaysi sohada o'zingizni rivojlantirishni xohlaysiz?",
                "options": ["Akademik bilimlar", "Liderlik ko'nikmalari", "Ilmiy izlanish", "Sport", "Ijodiy yo'nalishlar"]
            },
            {
                "type": "choice",
                "question": "11. Universitetga kirish uchun sizga birinchi navbatda qanday yordam kerak?",
                "options": ["Testlarga tayyorgarlik", "Universitet tanlashda maslahat", "Hujjat topshirish tartibi", "Grant imkoniyatlari"]
            },
            {
                "type": "text",
                "question": "12. Kelajakdagi asosiy maqsadingiz nima? Qisqacha yozing:"
            },
            {
                "type": "choice",
                "question": "13. Agar chet elda o'qish imkoniyati yaratilsa, undan foydalanishga tayyormisiz?",
                "options": ["Ha, albatta", "Ehtimol", "Yo'q"]
            },
            {
                "type": "choice",
                "question": "14. Maktab maslahatchisi bilan suhbatlashishni xohlaysizmi?",
                "options": ["Ha, suhbatlashmoqchiman", "Hozircha yo'q"]
            },
            {
                "type": "text",
                "question": "15. Maktab hayotini yaxshilash uchun qanday taklifingiz bor?"
            }
        ]
    }
]

# ================= HOLATLAR (FSM) =================
class SurveyState(StatesGroup):
    waiting_fio = State()
    waiting_grade_num = State()
    waiting_grade_letter = State()
    answering = State()

class SupportState(StatesGroup):
    waiting_message = State()

# ================= BAZA ISHGA TUSHIRISH =================
async def init_db():
    async with aiosqlite.connect(DB_NAME) as db:
        await db.execute("""
            CREATE TABLE IF NOT EXISTS users (
                user_id INTEGER PRIMARY KEY,
                full_name TEXT,
                username TEXT,
                joined_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """)
        await db.execute("""
            CREATE TABLE IF NOT EXISTS answers (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER,
                full_name TEXT,
                grade TEXT,
                survey_title TEXT,
                question TEXT,
                answer TEXT,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """)
        await db.execute("""
            CREATE TABLE IF NOT EXISTS feedback_map (
                admin_msg_id INTEGER PRIMARY KEY,
                user_id INTEGER,
                user_name TEXT
            )
        """)
        await db.commit()

# ================= TUGMALAR =================
def get_main_keyboard(is_admin: bool = False):
    keyboard = [
        [KeyboardButton(text="📝 So'rovnomani to'ldirish")],
        [KeyboardButton(text="✉️ Maslahatchiga murojaat")],
        [KeyboardButton(text="🎭 7 ta rivojlanish klubi"), KeyboardButton(text="👥 O'quvchilar kengashi")],
        [KeyboardButton(text="👨‍💼 Maktab maslahatchisi")]
    ]
    if is_admin:
        keyboard.append([KeyboardButton(text="📊 Excel hisobotni yuklash")])
    return ReplyKeyboardMarkup(keyboard=keyboard, resize_keyboard=True)

def get_cancel_keyboard():
    return ReplyKeyboardMarkup(
        keyboard=[[KeyboardButton(text="❌ Bekor qilish")]],
        resize_keyboard=True
    )

def get_grade_num_keyboard():
    keyboard = [
        [InlineKeyboardButton(text="5-sinf", callback_data="num_5"), InlineKeyboardButton(text="6-sinf", callback_data="num_6")],
        [InlineKeyboardButton(text="7-sinf", callback_data="num_7"), InlineKeyboardButton(text="8-sinf", callback_data="num_8")],
        [InlineKeyboardButton(text="9-sinf", callback_data="num_9"), InlineKeyboardButton(text="10-sinf", callback_data="num_10")],
        [InlineKeyboardButton(text="11-sinf", callback_data="num_11")],
        [InlineKeyboardButton(text="❌ Bekor qilish", callback_data="cancel_survey")]
    ]
    return InlineKeyboardMarkup(inline_keyboard=keyboard)

def get_grade_letter_keyboard():
    letters = ["A", "B", "V", "G", "D"]
    row = [InlineKeyboardButton(text=l, callback_data=f"let_{l}") for l in letters]
    keyboard = [row, [InlineKeyboardButton(text="❌ Bekor qilish", callback_data="cancel_survey")]]
    return InlineKeyboardMarkup(inline_keyboard=keyboard)

# ================= BOT INITIALIZATION =================
bot = Bot(
    token=BOT_TOKEN,
    default=DefaultBotProperties(parse_mode=ParseMode.HTML)
)
dp = Dispatcher()

# ================= ENG BIRINCHI: BEKOR QILISH HANDLERLARI =================
@dp.message(F.text == "❌ Bekor qilish")
async def cancel_any_action(message: Message, state: FSMContext):
    await state.clear()
    is_admin = (message.from_user.id == ADMIN_ID)
    await message.answer(
        "Jarayon bekor qilindi va asosiy menyuga qaytdingiz.",
        reply_markup=get_main_keyboard(is_admin)
    )

@dp.callback_query(F.data == "cancel_survey")
async def cancel_survey_callback(callback: CallbackQuery, state: FSMContext):
    await state.clear()
    is_admin = (callback.from_user.id == ADMIN_ID)
    await callback.message.delete()
    await callback.message.answer(
        "So'rovnoma bekor qilindi va asosiy menyuga qaytdingiz.",
        reply_markup=get_main_keyboard(is_admin)
    )
    await callback.answer()

@dp.message(CommandStart())
async def cmd_start(message: Message, state: FSMContext):
    await state.clear()
    user = message.from_user
    
    async with aiosqlite.connect(DB_NAME) as db:
        await db.execute(
            "INSERT OR REPLACE INTO users (user_id, full_name, username) VALUES (?, ?, ?)",
            (user.id, user.full_name, user.username or "")
        )
        await db.commit()

    is_admin = (user.id == ADMIN_ID)
    await message.answer(
        f"Assalomu alaykum, <b>{user.full_name}</b>!\n\n"
        "🏫 <b>Maktab maslahatchisining rasmiy portaliga xush kelibsiz.</b>\n"
        "Ushbu bot orqali siz o'z qiziqishlaringiz bo'yicha so'rovnomalarda qatnashishingiz, "
        "maktab klublari va kengash faoliyati bilan tanishishingiz hamda "
        "maslahatchiga to'g'ridan-to'g'ri murojaat yo'llashingiz mumkin.\n\n"
        "Kerakli bo'limni tanlang:",
        reply_markup=get_main_keyboard(is_admin)
    )

# ================= ASOSIY BO'LIMLAR =================
@dp.message(F.text == "👨‍💼 Maktab maslahatchisi")
async def handle_counselor(message: Message):
    await message.answer(ABOUT_COUNSELOR)

@dp.message(F.text == "🎭 7 ta rivojlanish klubi")
async def handle_clubs(message: Message):
    await message.answer(CLUBS_INFO)

@dp.message(F.text == "👥 O'quvchilar kengashi")
async def handle_council(message: Message):
    await message.answer(COUNCIL_INFO)

# ================= ADMIN BUYRUQLARI =================
@dp.message(F.chat.id == ADMIN_ID, Command("post"))
async def admin_broadcast_message(message: Message):
    text = message.text.replace("/post", "", 1).strip()
    if not text:
        await message.answer(
            "⚠️ <b>Matn kiritilmadi!</b>\n"
            "Foydalanish tartibi:\n"
            "<code>/post Barcha o'quvchilar diqqatiga! Ertaga tadbir bo'ladi.</code>"
        )
        return

    async with aiosqlite.connect(DB_NAME) as db:
        async with db.execute("SELECT user_id FROM users") as cursor:
            users = await cursor.fetchall()

    if not users:
        await message.answer("Bazada hali hech qanday foydalanuvchi yo'q.")
        return

    wait_msg = await message.answer(f"⏳ Xabar {len(users)} ta foydalanuvchiga yuborilmoqda...")
    success = 0
    failed = 0

    for (uid,) in users:
        try:
            await bot.send_message(
                chat_id=uid,
                text=f"📢 <b>Maktab maslahatchisidan umumiy e'lon:</b>\n"
                     f"━━━━━━━━━━━━━━━━━━━━\n"
                     f"{text}"
            )
            success += 1
            await asyncio.sleep(0.05)
        except Exception:
            failed += 1

    await wait_msg.delete()
    await message.answer(
        f"✅ <b>E'lon tarqatildi!</b>\n\n"
        f"• Muvaffaqiyatli yetib bordi: <b>{success} ta</b>\n"
        f"• Yetkazilmadi (bloklaganlar): <b>{failed} ta</b>"
    )

@dp.message(F.chat.id == ADMIN_ID, Command("sms"))
async def admin_direct_sms(message: Message):
    parts = message.text.split(maxsplit=2)
    if len(parts) < 3:
        await message.answer(
            "⚠️ <b>Noto'g'ri format!</b>\n"
            "Foydalanish tartibi:\n"
            "<code>/sms 123456789 Xonamga kirib o'tgin</code>"
        )
        return

    target_id_str, sms_text = parts[1], parts[2]
    try:
        target_id = int(target_id_str)
    except ValueError:
        await message.answer("❌ ID raqami faqat sonlardan iborat bo'lishi kerak!")
        return

    try:
        await bot.send_message(
            chat_id=target_id,
            text=f"👨‍💼 <b>Maktab maslahatchisidan shaxsiy xabar:</b>\n"
                 f"━━━━━━━━━━━━━━━━━━━━\n"
                 f"{sms_text}"
        )
        await message.answer(f"✅ Xabar <code>{target_id}</code> raqamli o'quvchiga muvaffaqiyatli yetkazildi!")
    except Exception:
        await message.answer("❌ Xabarni yuborib bo'lmadi (o'quvchi botni to'xtatgan bo'lishi mumkin).")

# ================= MUROJAAT BO'LIMI =================
@dp.message(F.text == "✉️ Maslahatchiga murojaat")
async def ask_counselor_start(message: Message, state: FSMContext):
    await message.answer(
        "✍️ <b>Maslahatchiga yo'llamoqchi bo'lgan savol yoki murojaatingizni yozing:</b>\n\n"
        "<i>Siz yozgan xabar maslahatchiga yetkaziladi va bot orqali javob qaytariladi. "
        "Bekor qilish uchun pastdagi tugmani bosing.</i>",
        reply_markup=get_cancel_keyboard()
    )
    await state.set_state(SupportState.waiting_message)

@dp.message(SupportState.waiting_message)
async def process_user_murojaat(message: Message, state: FSMContext):
    u_name = message.from_user.full_name
    u_id = message.from_user.id
    u_text = message.text

    try:
        admin_forward = await bot.send_message(
            chat_id=ADMIN_ID,
            text=f"📬 <b>Yangi murojaat keldi:</b>\n"
                 f"👤 <b>Foydalanuvchi:</b> {u_name} (ID: <code>{u_id}</code>)\n"
                 f"━━━━━━━━━━━━━━━━━━━━\n"
                 f"💬 <b>Xabar matni:</b>\n{u_text}\n\n"
                 f"<i>Ushbu xabarga Telegram'ning 'Ответить' (Reply) tugmasi bilan javob bering yoki <code>/sms {u_id} Javobingiz</code> qiling.</i>"
        )

        async with aiosqlite.connect(DB_NAME) as db:
            await db.execute(
                "INSERT INTO feedback_map (admin_msg_id, user_id, user_name) VALUES (?, ?, ?)",
                (admin_forward.message_id, u_id, u_name)
            )
            await db.commit()

        is_admin = (u_id == ADMIN_ID)
        await message.answer(
            "✅ <b>Xabaringiz maktab maslahatchisiga yetkazildi!</b>\n"
            "Tez orada bot orqali sizga javob qaytariladi.",
            reply_markup=get_main_keyboard(is_admin)
        )
    except Exception:
        await message.answer("Xabarni yuborishda xatolik yuz berdi. Keyinroq urinib ko'ring.")
    
    await state.clear()

# ================= ADMIN REPLY JAVOB =================
@dp.message(F.chat.id == ADMIN_ID, F.reply_to_message)
async def admin_reply_to_user(message: Message):
    replied_id = message.reply_to_message.message_id

    async with aiosqlite.connect(DB_NAME) as db:
        async with db.execute("SELECT user_id, user_name FROM feedback_map WHERE admin_msg_id = ?", (replied_id,)) as cursor:
            row = await cursor.fetchone()

    if row:
        target_user_id, target_name = row
        try:
            await bot.send_message(
                chat_id=target_user_id,
                text=f"👨‍💼 <b>Maktab maslahatchisidan javob keldi:</b>\n"
                     f"━━━━━━━━━━━━━━━━━━━━\n"
                     f"{message.text}"
            )
            await message.reply(f"✅ Javobingiz <b>{target_name}</b>ga muvaffaqiyatli yuborildi!")
        except Exception:
            await message.reply("❌ Javobni yuborib bo'lmadi (foydalanuvchi botni bloklagan bo'lishi mumkin).")

# ================= SO'ROVNOMA BOSQICHLARI =================
@dp.message(F.text == "📝 So'rovnomani to'ldirish")
async def start_survey(message: Message, state: FSMContext):
    await message.answer(
        "📝 <b>So'rovnomani to'ldirish uchun avval Ism va Familiyangizni kiriting:</b>\n"
        "<i>(Masalan: Karimov Jasur)</i>\n\n"
        "<i>Jarayonni to'xtatmoqchi bo'lsangiz, '❌ Bekor qilish' tugmasini bosing.</i>",
        reply_markup=get_cancel_keyboard()
    )
    await state.set_state(SurveyState.waiting_fio)

@dp.message(SurveyState.waiting_fio)
async def process_fio(message: Message, state: FSMContext):
    fio = message.text.strip()
    await state.update_data(fio=fio)
    await message.answer(
        f"Rahmat, <b>{fio}</b>!\nEndi sinfingizni tanlang:",
        reply_markup=get_grade_num_keyboard()
    )
    await state.set_state(SurveyState.waiting_grade_num)

@dp.callback_query(SurveyState.waiting_grade_num, F.data.startswith("num_"))
async def process_grade_num(callback: CallbackQuery, state: FSMContext):
    grade_num = int(callback.data.split("_")[1])
    await state.update_data(grade_num=grade_num)
    await callback.message.edit_text(
        f"Tanlandi: <b>{grade_num}-sinf</b>\nSinfingiz harfini tanlang:",
        reply_markup=get_grade_letter_keyboard()
    )
    await state.set_state(SurveyState.waiting_grade_letter)
    await callback.answer()

@dp.callback_query(SurveyState.waiting_grade_letter, F.data.startswith("let_"))
async def process_grade_letter(callback: CallbackQuery, state: FSMContext):
    grade_letter = callback.data.split("_")[1]
    data = await state.get_data()
    grade_num = data["grade_num"]
    full_grade = f"{grade_num}-{grade_letter}"
    await state.update_data(full_grade=full_grade)

    selected_survey = None
    for s in SURVEYS:
        if grade_num in s["grades"]:
            selected_survey = s
            break

    if not selected_survey:
        is_admin = (callback.from_user.id == ADMIN_ID)
        await callback.message.delete()
        await callback.message.answer(
            f"{full_grade}-sinf uchun ayni vaqtda faol so'rovnoma mavjud emas.",
            reply_markup=get_main_keyboard(is_admin)
        )
        await state.clear()
        return

    await state.update_data(
        survey=selected_survey,
        q_idx=0,
        answers=[]
    )
    await callback.message.delete()
    await send_next_question(callback.message, state, is_first=True)
    await callback.answer()

async def send_next_question(message: Message, state: FSMContext, is_first: bool = False):
    data = await state.get_data()
    survey = data["survey"]
    q_idx = data["q_idx"]
    total = len(survey["questions"])

    q_data = survey["questions"][q_idx]
    q_text = q_data["question"]
    q_type = q_data["type"]

    progress_bar = f"📊 <b>Savol: {q_idx + 1} / {total}</b>\n"
    full_text = f"{progress_bar}━━━━━━━━━━━━━━━━━━━━\n<b>{q_text}</b>"

    if q_type == "choice":
        keyboard = []
        for opt in q_data["options"]:
            cb_data = f"ans_{opt[:25]}"
            keyboard.append([InlineKeyboardButton(text=opt, callback_data=cb_data)])
        keyboard.append([InlineKeyboardButton(text="❌ So'rovnomani to'xtatish", callback_data="cancel_survey")])
        
        reply_markup = InlineKeyboardMarkup(inline_keyboard=keyboard)
        await message.answer(full_text, reply_markup=reply_markup)
        await state.set_state(SurveyState.answering)
    else:
        full_text += "\n\n✍️ <i>Javobingizni xabar sifatida yozib yuboring (yoki pastdagi tugma orqali to'xtating):</i>"
        await message.answer(full_text, reply_markup=get_cancel_keyboard())
        await state.set_state(SurveyState.answering)

@dp.callback_query(SurveyState.answering, F.data.startswith("ans_"))
async def handle_choice_answer(callback: CallbackQuery, state: FSMContext):
    chosen_short = callback.data.replace("ans_", "")
    data = await state.get_data()
    survey = data["survey"]
    q_idx = data["q_idx"]
    q_data = survey["questions"][q_idx]

    full_answer = next((opt for opt in q_data["options"] if opt.startswith(chosen_short)), chosen_short)
    await save_and_proceed(callback.message, state, q_data["question"], full_answer, callback.from_user.id)
    await callback.answer()

@dp.message(SurveyState.answering)
async def handle_text_answer(message: Message, state: FSMContext):
    data = await state.get_data()
    survey = data["survey"]
    q_idx = data["q_idx"]
    q_data = survey["questions"][q_idx]

    await save_and_proceed(message, state, q_data["question"], message.text.strip(), message.from_user.id)

async def save_and_proceed(message: Message, state: FSMContext, question: str, answer: str, user_id: int):
    data = await state.get_data()
    answers = data["answers"]
    answers.append((question, answer))
    q_idx = data["q_idx"] + 1
    survey = data["survey"]

    if q_idx < len(survey["questions"]):
        await state.update_data(q_idx=q_idx, answers=answers)
        await send_next_question(message, state)
    else:
        fio = data["fio"]
        grade = data["full_grade"]
        survey_title = survey["title"]

        async with aiosqlite.connect(DB_NAME) as db:
            for q, a in answers:
                await db.execute(
                    "INSERT INTO answers (user_id, full_name, grade, survey_title, question, answer) VALUES (?, ?, ?, ?, ?, ?)",
                    (user_id, fio, grade, survey_title, q, a)
                )
            await db.commit()

        await state.clear()
        is_admin = (user_id == ADMIN_ID)
        await message.answer(
            f"✅ <b>Katta rahmat, {fio}!</b>\n\n"
            "Siz so'rovnomani to'liq yakunladingiz. "
            "Barcha javoblaringiz maslahatchiga muvaffaqiyatli yetkazildi.",
            reply_markup=get_main_keyboard(is_admin)
        )

# ================= FORMATLANGAN EXCEL HISOBOT =================
@dp.message(F.text == "📊 Excel hisobotni yuklash")
async def generate_excel_report(message: Message):
    if message.from_user.id != ADMIN_ID:
        await message.answer("Ushbu bo'lim faqat maktab maslahatchisi uchun ochiq!")
        return

    wait_msg = await message.answer("⏳ So'rovnoma natijalari Excel faylga tayyorlanmoqda...")

    async with aiosqlite.connect(DB_NAME) as db:
        query = "SELECT user_id, full_name, grade, survey_title, question, answer, id FROM answers ORDER BY id ASC"
        async with db.execute(query) as cursor:
            rows = await cursor.fetchall()

    if not rows:
        await wait_msg.edit_text("Bazada hozircha to'ldirilgan so'rovnomalar mavjud emas.")
        return

    raw_df = pd.DataFrame(rows, columns=["user_id", "full_name", "grade", "survey_title", "question", "answer", "id"])
    file_path = "Maktab_Sorovnoma_Hisoboti.xlsx"

    with pd.ExcelWriter(file_path, engine='openpyxl') as writer:
        unique_grades = sorted(raw_df["grade"].unique())
        
        for grade in unique_grades:
            grade_data = raw_df[raw_df["grade"] == grade]
            students = grade_data[["user_id", "full_name"]].drop_duplicates().values.tolist()
            
            questions = []
            for q in grade_data["question"]:
                if q not in questions:
                    questions.append(q)
            
            rows_list = []
            for idx, (u_id, fio) in enumerate(students, start=1):
                row = {
                    "T/r": idx,
                    "O'quvchining F.I.SH": fio,
                    "Telegram ID": u_id
                }
                st_answers = grade_data[(grade_data["user_id"] == u_id) & (grade_data["full_name"] == fio)]
                for q in questions:
                    ans_series = st_answers[st_answers["question"] == q]["answer"]
                    row[q] = ans_series.iloc[-1] if not ans_series.empty else "-"
                
                rows_list.append(row)
            
            sheet_df = pd.DataFrame(rows_list)
            sheet_name = str(grade).replace("/", "-")[:30]
            sheet_df.to_excel(writer, sheet_name=sheet_name, index=False)
            
            worksheet = writer.sheets[sheet_name]
            
            header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            border_style = Border(
                left=Side(style='thin', color='D9D9D9'),
                right=Side(style='thin', color='D9D9D9'),
                top=Side(style='thin', color='D9D9D9'),
                bottom=Side(style='thin', color='D9D9D9')
            )

            for col_num in range(1, len(sheet_df.columns) + 1):
                cell = worksheet.cell(row=1, column=col_num)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = border_style

            worksheet.row_dimensions[1].height = 40

            zebra_fill = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
            for row_idx in range(2, len(rows_list) + 2):
                is_zebra = (row_idx % 2 == 1)
                for col_idx in range(1, len(sheet_df.columns) + 1):
                    cell = worksheet.cell(row=row_idx, column=col_idx)
                    cell.border = border_style
                    if is_zebra:
                        cell.fill = zebra_fill
                    
                    if col_idx in [1, 3]:
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                    elif col_idx == 2:
                        cell.alignment = Alignment(horizontal="left", vertical="center")
                    else:
                        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

            for col in worksheet.columns:
                col_letter = get_column_letter(col[0].column)
                if col[0].column == 1:
                    worksheet.column_dimensions[col_letter].width = 8
                elif col[0].column == 2:
                    worksheet.column_dimensions[col_letter].width = 30
                elif col[0].column == 3:
                    worksheet.column_dimensions[col_letter].width = 16
                else:
                    worksheet.column_dimensions[col_letter].width = 32

    await wait_msg.delete()
    doc = FSInputFile(file_path)
    await message.answer_document(
        doc,
        caption="📊 <b>Maktab so'rovnomalari bo'yicha shakllantirilgan Excel hisoboti</b>\n"
                "<i>(Har bir o'quvchining shaxsiy Telegram ID raqami ham kiritilgan)</i>"
    )

    if os.path.exists(file_path):
        os.remove(file_path)

# ================= RENDER UCHUN MINI VEB-SERVER =================
async def handle_ping(request):
    return web.Response(text="Bot 24/7 faol ishlamoqda!")

async def start_web_server():
    app = web.Application()
    app.router.add_get('/', handle_ping)
    runner = web.AppRunner(app)
    await runner.setup()
    port = int(os.environ.get("PORT", 10000))
    site = web.TCPSite(runner, '0.0.0.0', port)
    await site.start()

# ================= ISHGA TUSHIRISH =================
async def main():
    await init_db()
    await start_web_server()
    await bot.delete_webhook(drop_pending_updates=True)
    print("Bot va veb-server muvaffaqiyatli ishga tushdi!")
    await dp.start_polling(bot)

if __name__ == "__main__":
    asyncio.run(main())
